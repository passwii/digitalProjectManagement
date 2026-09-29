"""资产导入 / 导出：用户看到的是中文，数据库仍是英文编码。"""
from datetime import date
import csv
import io
from urllib.parse import unquote

import pytest
from openpyxl import Workbook, load_workbook

from api.models import Dict, EquipmentAsset
from test_mvp import login

ROOT = "/api/v1/equipment-assets"


@pytest.fixture(autouse=True)
def categories(db):
    with db() as session:
        for code, label in (("computer", "电脑"), ("monitor", "显示器"), ("printer", "打印机"),
                            ("network", "网络设备"), ("instrument", "工具仪器"), ("other", "其他")):
            session.add(Dict(type="equipment_category", code=code, label=label))
        session.commit()


def _upload(client, headers, filename, payload):
    return client.post(
        f"{ROOT}/import", headers=headers, data={"batch_name": f"导入 {filename}"},
        files={"file": (filename, payload)},
    )


def _filename(response) -> str:
    disposition = response.headers["content-disposition"]
    return unquote(disposition.split("filename*=UTF-8''")[-1].strip('"'))


def test_import_accepts_chinese_headers_and_local_values(client, db):
    admin = login(client, "admin")
    # 字段带引号，模拟用户在 Excel 里输入 ¥7,200.50 后另存为 CSV 的真实结果
    payload = ('"设备名称","设备分类","序列号","品牌","型号","存放位置","采购日期","采购金额","备注"\n'
               '"中文测试电脑","电脑","ZH-1","联想","ThinkPad","总部 5 楼","2026/3/8","¥7,200.50","备注一行"\n').encode()
    response = _upload(client, admin, "台账.csv", payload)
    assert response.status_code == 200, response.text
    assert response.json()["created"] == 1
    with db() as session:
        asset = session.query(EquipmentAsset).one()
        # 英文编码进库，中文只停留在文件边界
        assert asset.category_code == "computer"
        assert asset.purchase_date == date(2026, 3, 8)
        assert float(asset.purchase_price) == 7200.50
        assert asset.note == "备注一行"


def test_import_accepts_alternative_chinese_headers(client):
    admin = login(client, "admin")
    payload = "名称,分类,购置日期,金额\n别名写法,显示器,2026年1月5日,1200\n".encode()
    response = _upload(client, admin, "别名.csv", payload)
    assert response.status_code == 200, response.text
    assert response.json()["created"] == 1


def test_import_missing_column_reports_chinese_column_name(client):
    admin = login(client, "admin")
    response = _upload(client, admin, "缺列.csv", "设备名称,备注\n只有名称,无分类\n".encode())
    assert response.status_code == 400, response.text
    assert "设备分类" in response.json()["detail"]


def test_import_error_rows_are_chinese_with_row_numbers(client, db):
    admin = login(client, "admin")
    payload = ("设备名称,设备分类,采购金额\n"
               "正常设备,电脑,100\n"
               "未知分类设备,不存在的分类,200\n"
               "金额写错设备,电脑,不是数字\n").encode()
    response = _upload(client, admin, "错误.csv", payload)
    assert response.status_code == 400, response.text
    detail = response.json()["detail"]
    assert "共 2 行导入失败" in detail["message"]
    assert detail["errors"] == [
        {"row": 3, "message": "设备分类「不存在的分类」不存在，请使用已有分类名称或编码"},
        {"row": 4, "message": "采购金额格式错误，请填数字"},
    ]
    with db() as session:
        assert session.query(EquipmentAsset).count() == 0


def test_import_reports_ignored_export_only_and_unknown_columns(client):
    admin = login(client, "admin")
    payload = ("资产编号,设备名称,设备分类,设备状态,使用人,神秘列\n"
               "EQ-2026-XYZ,带额外列的设备,电脑,可用,张三,随便写的\n").encode()
    response = _upload(client, admin, "回传.csv", payload)
    assert response.status_code == 200, response.text
    ignored = response.json()["ignored_columns"]
    assert ignored["未识别"] == ["神秘列"]
    assert set(ignored["仅导出用(导入时忽略)"]) == {"资产编号", "设备状态", "使用人"}


def test_export_uses_chinese_headers_and_chinese_values(client):
    admin, pm = login(client, "admin"), login(client, "pm")
    created = client.post(f"{ROOT}/inbound", headers=admin, json={
        "name": "导出测试批次", "source": "采购单 1", "items": [{"name": "导出测试电脑", "category_code": "computer",
                                        "serial_number": "EXP-1", "purchase_price": 1234.5}],
    })
    batch_id = created.json()["batch_id"]
    asset_id = created.json()["items"][0]["id"]
    assert client.post(f"{ROOT}/{asset_id}/borrow", headers=pm, json={"due_date": "2099-01-01"}).status_code == 200

    response = client.get(f"{ROOT}/export?format=csv", headers=admin)
    assert response.status_code == 200
    rows = list(csv.DictReader(io.StringIO(response.content.decode("utf-8-sig"))))
    row = rows[0]
    # 表头全中文
    assert "设备名称" in row and "category_code" not in row
    # 取值也是中文，不是 available / computer / user id
    assert row["设备分类"] == "电脑"
    assert row["设备状态"] == "借用中"
    assert row["使用人"] == "pm"
    assert row["归还日期"] == "2099-01-01"
    assert "导出测试批次" in row["所属批次"] and f"#{batch_id}" in row["所属批次"]
    assert float(row["采购金额"]) == 1234.5
    assert row["设备名称"] == "导出测试电脑"
    assert _filename(response) == "设备台账.csv"


def test_export_xlsx_keeps_numbers_numeric_and_chinese_sheet(client):
    admin = login(client, "admin")
    client.post(f"{ROOT}", headers=admin, json={"name": "数值设备", "category_code": "computer",
                                               "purchase_price": 99.99, "purchase_date": "2026-02-02"})
    response = client.get(f"{ROOT}/export?format=xlsx", headers=admin)
    sheet = load_workbook(io.BytesIO(response.content), data_only=True).active
    assert sheet.title == "设备台账"
    headers = [cell.value for cell in sheet[1]]
    assert headers[:3] == ["资产编号", "设备名称", "设备分类"]
    amount = sheet.cell(2, headers.index("采购金额") + 1)
    # 金额保持数值型，用户在 Excel 里可以直接求和
    assert amount.value == 99.99
    assert sheet.cell(2, headers.index("设备分类") + 1).value == "电脑"
    assert sheet.cell(2, headers.index("采购日期") + 1).value == "2026-02-02"


def test_export_then_reimport_round_trip_keeps_chinese_category(client, db):
    admin = login(client, "admin")
    client.post(f"{ROOT}", headers=admin, json={"name": "往返设备", "category_code": "printer",
                                               "serial_number": "SN-AAA"})
    exported = client.get(f"{ROOT}/export?format=csv", headers=admin)
    # 导出的文件改个名和序列号回传，应当被接受（分类列是中文名称，不是 printer）
    payload = exported.content.replace("往返设备".encode(), "往返设备二".encode()) \
                             .replace("SN-AAA".encode(), "SN-BBB".encode())
    response = _upload(client, admin, "回传导出.csv", payload)
    assert response.status_code == 200, response.text
    assert response.json()["created"] == 1
    with db() as session:
        assert session.query(EquipmentAsset).filter_by(name="往返设备二").one().category_code == "printer"


def test_import_template_is_chinese_and_lists_categories(client):
    admin, pm = login(client, "admin"), login(client, "pm")
    assert client.get(f"{ROOT}/import-template?format=csv", headers=pm).status_code == 403

    csv_response = client.get(f"{ROOT}/import-template?format=csv", headers=admin)
    assert csv_response.status_code == 200
    rows = list(csv.reader(io.StringIO(csv_response.content.decode("utf-8-sig"))))
    assert rows[0] == ["设备名称", "设备分类", "序列号", "品牌", "型号", "存放位置", "采购日期", "采购金额", "备注"]

    xlsx_response = client.get(f"{ROOT}/import-template?format=xlsx", headers=admin)
    book = load_workbook(io.BytesIO(xlsx_response.content))
    assert [c.value for c in book["台账导入"][1]] == rows[0]
    guide = [[c.value for c in row] for row in book["填写说明"].iter_rows()]
    assert guide[0] == ["列名", "是否必填", "填写说明"]
    assert "电脑" in guide[-1][2]
    assert _filename(xlsx_response) == "设备台账导入模板.xlsx"
