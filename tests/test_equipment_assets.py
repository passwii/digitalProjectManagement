"""公司设备资产：权限、生命周期和批量导入的接口契约。"""
from datetime import date
import csv
import io
import json

import pytest
from openpyxl import Workbook
from sqlalchemy import text

from api.core.timeutil import app_today
from api.models import AuditLog, Dict, EquipmentAsset, EquipmentEvent, EquipmentBatch, User
from test_mvp import login


ROOT = "/api/v1/equipment-assets"

# 与 0027/0030 迁移完全一致的可变体。单元测试用 create_all 建表、**不跑迁移**，
# 所以流水不可变触发器在这里必须手工装上——否则「删流水被触发器拦截导致生产 500」
# 这类缺陷在测试里永远是绿的（#99 删除功能就是这么漏到线上的）。
IMMUTABLE_TRIGGER_FN = """
CREATE OR REPLACE FUNCTION prevent_equipment_event_change()
RETURNS trigger
LANGUAGE plpgsql
AS $function$
BEGIN
    IF COALESCE(current_setting('app.allow_event_purge', true), 'off') = 'on' THEN
        RETURN OLD;
    END IF;
    RAISE EXCEPTION 'equipment event history is immutable';
END; $function$
"""
IMMUTABLE_TRIGGER = (
    "CREATE TRIGGER equipment_events_immutable BEFORE UPDATE OR DELETE ON equipment_events "
    "FOR EACH ROW EXECUTE FUNCTION prevent_equipment_event_change()"
)


@pytest.fixture(autouse=True)
def categories(db):
    """分类字典 + 流水不可变触发器。

    `db` 是 sessionmaker，`with db()` 拿到的会话与被测代码（TestClient 走
    dependency_overrides）指向同一个临时 schema，所以触发器对接口同样生效。
    """
    with db() as session:
        for code, label in (("computer", "电脑"), ("monitor", "显示器"), ("printer", "打印机"),
                            ("network", "网络设备"), ("instrument", "工具仪器"), ("other", "其他")):
            session.add(Dict(type="equipment_category", code=code, label=label))
        session.commit()
        session.execute(text(IMMUTABLE_TRIGGER_FN))
        session.execute(text(IMMUTABLE_TRIGGER))
        session.commit()


@pytest.fixture()
def locked_events(client, db):
    """反向确认：触发器真的在拦截（否则删除相关测试是假绿）。

    0027 的触发器是 BEFORE UPDATE OR DELETE + FOR EACH ROW，所以必须先有真实流水行，
    0 行的 DELETE 不会触发。探针资产与流水随临时 schema 一起丢弃，删不掉也无妨。
    """
    probe = _create(client, login(client, "admin"), name="触发器探针")
    with db() as session:
        eid = (session.query(EquipmentEvent)
               .filter_by(asset_id=probe["id"]).order_by(EquipmentEvent.id.desc()).first().id)
        for sql in ("DELETE FROM equipment_events WHERE id = :i",
                    "UPDATE equipment_events SET note = 'x' WHERE id = :i"):
            with pytest.raises(Exception, match="equipment event history is immutable"):
                session.execute(text(sql), {"i": eid})
            session.rollback()


def _create(client, headers, name="笔记本", **extra):
    response = client.post(ROOT, headers=headers, json={"name": name, "category_code": "computer", **extra})
    assert response.status_code == 200, response.text
    return response.json()


def test_permissions_and_lifecycle(client, db):
    admin, pm, outsider = (login(client, who) for who in ("admin", "pm", "outsider"))
    assert client.post(ROOT, headers=pm, json={"name": "电脑", "category_code": "computer"}).status_code == 403
    assert client.get(ROOT, headers=pm).status_code == 200
    asset = _create(client, admin)
    aid = asset["id"]
    assert asset["code"].startswith("EQ-")
    assert client.post(f"{ROOT}/{aid}/borrow", headers=pm, json={"due_date": "2099-01-01"}).status_code == 200
    assert client.post(f"{ROOT}/{aid}/return", headers=outsider).status_code == 403
    assert client.post(f"{ROOT}/{aid}/return", headers=pm).status_code == 200
    assert client.post(f"{ROOT}/{aid}/claim", headers=pm).status_code == 200
    assert client.post(f"{ROOT}/{aid}/return", headers=pm).status_code == 403
    assert client.post(f"{ROOT}/{aid}/return", headers=admin).status_code == 200
    assert client.post(f"{ROOT}/{aid}/maintenance", headers=pm).status_code == 403
    assert client.post(f"{ROOT}/{aid}/maintenance", headers=admin).status_code == 200
    assert client.post(f"{ROOT}/{aid}/scrap", headers=admin).status_code == 200
    assert client.post(f"{ROOT}/{aid}/borrow", headers=pm).status_code == 400
    events = client.get(f"{ROOT}/{aid}/events", headers=pm)
    assert [item["action"] for item in events.json()["items"]] == [
        "create", "borrow", "return", "claim", "return", "maintenance", "scrap"
    ]
    with db() as session:
        assert session.query(EquipmentEvent).filter_by(asset_id=aid).count() == 7
        assert session.query(AuditLog).filter_by(resource="equipment_assets", resource_id=str(aid)).count() == 7


def test_admin_can_permanently_delete_only_correction_only_assets(client, db):
    admin, pm = login(client, "admin"), login(client, "pm")
    asset = _create(client, admin, name="重复录入电脑", serial_number="DELETE-ME-1")
    aid = asset["id"]

    assert client.get(f"{ROOT}/{aid}/delete-check", headers=pm).status_code == 403
    assert client.request("DELETE", f"{ROOT}/{aid}", headers=pm,
                          json={"confirm_code": asset["code"], "reason": "重复录入"}).status_code == 403

    check = client.get(f"{ROOT}/{aid}/delete-check", headers=admin)
    assert check.status_code == 200, check.text
    assert check.json() == {
        "allowed": True, "asset_id": aid, "code": asset["code"], "reason": None,
        "event_count": 1, "protected_actions": [],
    }

    updated = client.patch(f"{ROOT}/{aid}", headers=admin, json={
        "name": "重复录入电脑（已核对）", "category_code": "computer", "serial_number": "DELETE-ME-1",
    })
    assert updated.status_code == 200, updated.text
    check = client.get(f"{ROOT}/{aid}/delete-check", headers=admin).json()
    assert check["allowed"] is True
    assert check["event_count"] == 2

    assert client.request("DELETE", f"{ROOT}/{aid}", headers=admin,
                          json={"confirm_code": "WRONG-CODE", "reason": "重复录入"}).status_code == 400
    assert client.request("DELETE", f"{ROOT}/{aid}", headers=admin,
                          json={"confirm_code": asset["code"], "reason": "   "}).status_code == 400

    deleted = client.request("DELETE", f"{ROOT}/{aid}", headers=admin,
                             json={"confirm_code": asset["code"], "reason": "重复录入"})
    assert deleted.status_code == 200, deleted.text
    assert deleted.json() == {"ok": True, "id": aid, "code": asset["code"]}

    with db() as session:
        assert session.get(EquipmentAsset, aid) is None
        assert session.query(EquipmentEvent).filter_by(asset_id=aid).count() == 0
        audit = session.query(AuditLog).filter_by(
            action="delete_equipment", resource="equipment_assets", resource_id=str(aid)
        ).one()
        detail = json.loads(audit.detail)
        assert detail["code"] == asset["code"]
        assert detail["name"] == "重复录入电脑（已核对）"
        assert detail["reason"] == "重复录入"
        assert detail["event_count"] == 2


def test_acceptance_does_not_block_delete_and_batch_is_preserved(client, db):
    admin = login(client, "admin")
    inbound = client.post(f"{ROOT}/inbound", headers=admin, json={
        "name": "验收后删除批次", "supplier": "录入测试供应商", "acceptance_status": "pending_acceptance",
        "items": [{"name": "误录显示器", "category_code": "monitor"}],
    })
    assert inbound.status_code == 200, inbound.text
    batch_id = inbound.json()["batch_id"]
    asset = inbound.json()["items"][0]
    assert client.post(f"{ROOT}/batches/{batch_id}/accept", headers=admin).status_code == 200

    check = client.get(f"{ROOT}/{asset['id']}/delete-check", headers=admin)
    assert check.status_code == 200, check.text
    assert check.json()["allowed"] is True
    assert check.json()["event_count"] == 2

    deleted = client.request("DELETE", f"{ROOT}/{asset['id']}", headers=admin,
                             json={"confirm_code": asset["code"], "reason": "验收后发现重复录入"})
    assert deleted.status_code == 200, deleted.text
    with db() as session:
        assert session.get(EquipmentBatch, batch_id) is not None
        assert session.get(EquipmentAsset, asset["id"]) is None


def test_admin_can_force_delete_asset_with_business_lifecycle(client, db):
    admin, pm = login(client, "admin"), login(client, "pm")
    asset = _create(client, admin, name="已使用电脑")
    aid = asset["id"]
    assert client.post(f"{ROOT}/{aid}/borrow", headers=pm,
                       json={"due_date": "2099-01-01"}).status_code == 200
    assert client.post(f"{ROOT}/{aid}/return", headers=pm).status_code == 200

    # delete-check 仍如实回报"有受保护流水"，仅供前端展示警告，不再作为拦截闸门
    check = client.get(f"{ROOT}/{aid}/delete-check", headers=admin)
    assert check.status_code == 200, check.text
    assert check.json()["allowed"] is False
    assert check.json()["protected_actions"] == ["borrow", "return"]
    assert "借用" in check.json()["reason"]

    # 技术校验仍然生效：编号不匹配 / 原因缺失 依旧拒绝
    assert client.request("DELETE", f"{ROOT}/{aid}", headers=admin,
                          json={"confirm_code": "WRONG", "reason": "清理"}).status_code == 400
    assert client.request("DELETE", f"{ROOT}/{aid}", headers=admin,
                          json={"confirm_code": asset["code"], "reason": "   "}).status_code == 400
    with db() as session:
        assert session.get(EquipmentAsset, aid) is not None

    deleted = client.request("DELETE", f"{ROOT}/{aid}", headers=admin,
                             json={"confirm_code": asset["code"], "reason": "测试数据清理"})
    assert deleted.status_code == 200, deleted.text
    assert deleted.json() == {"ok": True, "id": aid, "code": asset["code"]}

    with db() as session:
        assert session.get(EquipmentAsset, aid) is None
        assert session.query(EquipmentEvent).filter_by(asset_id=aid).count() == 0
        audit = session.query(AuditLog).filter_by(
            action="delete_equipment", resource="equipment_assets", resource_id=str(aid)
        ).one()
        detail = json.loads(audit.detail)
        assert detail["force_delete"] is True
        assert detail["protected_actions"] == ["borrow", "return"]
        assert detail["reason"] == "测试数据清理"


def test_delete_survives_immutable_trigger_and_trigger_stays_closed_otherwise(client, db, locked_events):
    """回归：#99 的删除接口在生产会 500，因为 equipment_events 的不可变触发器拦截 DELETE。

    现在删除接口在事务内 SET LOCAL 打开 app.allow_event_purge 穿透开关，而触发器对
    其它写路径依然完全关闭。
    """
    admin, pm = login(client, "admin"), login(client, "pm")
    asset = _create(client, admin, name="带流水电脑")
    aid = asset["id"]
    assert client.post(f"{ROOT}/{aid}/borrow", headers=pm, json={"due_date": "2099-01-01"}).status_code == 200
    assert client.post(f"{ROOT}/{aid}/return", headers=pm).status_code == 200

    deleted = client.request("DELETE", f"{ROOT}/{aid}", headers=admin,
                             json={"confirm_code": asset["code"], "reason": "测试数据清理"})
    assert deleted.status_code == 200, deleted.text
    assert deleted.json() == {"ok": True, "id": aid, "code": asset["code"]}
    with db() as session:
        assert session.get(EquipmentAsset, aid) is None
        assert session.query(EquipmentEvent).filter_by(asset_id=aid).count() == 0

    # 关键不变量：删除之后，流水不可变依旧对普通写路径生效（事务级开关不泄漏）。
    keeper = _create(client, admin, name="留存设备")
    with db() as session:
        kid = keeper["id"]
        assert session.query(EquipmentEvent).filter_by(asset_id=kid).count() >= 1
        with pytest.raises(Exception, match="equipment event history is immutable"):
            session.execute(text("DELETE FROM equipment_events WHERE asset_id = :aid"), {"aid": kid})
        session.rollback()
        # UPDATE 同样被拦（0027 触发器是 BEFORE UPDATE OR DELETE）
        with pytest.raises(Exception, match="equipment event history is immutable"):
            session.execute(text("UPDATE equipment_events SET note = 'x' WHERE asset_id = :aid"), {"aid": kid})
        session.rollback()


def test_force_delete_of_clean_asset_is_not_flagged_in_audit(client, db, locked_events):
    admin = login(client, "admin")
    asset = _create(client, admin, name="重复录入电脑")
    deleted = client.request("DELETE", f"{ROOT}/{asset['id']}", headers=admin,
                             json={"confirm_code": asset["code"], "reason": "重复录入"})
    assert deleted.status_code == 200, deleted.text
    with db() as session:
        audit = session.query(AuditLog).filter_by(
            action="delete_equipment", resource="equipment_assets", resource_id=str(asset["id"])
        ).one()
        detail = json.loads(audit.detail)
        assert detail["force_delete"] is False
        assert detail["protected_actions"] == []


def test_batch_and_unique_codes(client, db):
    admin = login(client, "admin")
    response = client.post(f"{ROOT}/inbound", headers=admin, json={
        "name": "采购批次 A", "source": "采购单 1", "items": [{"name": "显示器", "category_code": "monitor"}, {"name": "显示器", "category_code": "monitor"}]
    })
    assert response.status_code == 200, response.text
    items = response.json()["items"]
    assert len(items) == 2 and len({item["code"] for item in items}) == 2
    assert all(item["batch_id"] == response.json()["batch_id"] for item in items)
    assert response.json()["acceptance_status"] == "accepted"
    assert response.json()["name"] == "采购批次 A"
    assert all(item["status"] == "available" for item in items)
    assert client.post(f"{ROOT}/inbound", headers=login(client, "pm"), json={"items": []}).status_code == 403


def test_inventory_inbound_outbound_tracks_source_quantity_and_shipment(client, db):
    admin = login(client, "admin")
    inbound = client.post(f"{ROOT}/inbound", headers=admin, json={
        "name": "库存网线入库",
        "inventory_managed": True,
        "items": [{
            "name": "Cat6 网线",
            "category_code": "network",
            "quantity": 12,
        }],
    })
    assert inbound.status_code == 200, inbound.text
    stock = inbound.json()["items"][0]
    assert stock["inventory_managed"] is True
    assert stock["quantity"] == 12
    assert stock["available_quantity"] == 12

    outbound = client.post(f"{ROOT}/{stock['id']}/outbound", headers=admin, json={
        "quantity": 5,
        "customer": "示例客户",
        "destination": "客户机房",
        "shipped_at": "2026-10-01T14:30:00+08:00",
        "note": "现场部署",
    })
    assert outbound.status_code == 200, outbound.text
    event = outbound.json()
    assert event["asset_id"] == stock["id"]
    assert event["batch_id"] == stock["batch_id"]
    assert event["quantity"] == 5
    assert event["customer"] == "示例客户"
    assert event["destination"] == "客户机房"
    assert event["shipped_at"] == "2026-10-01T14:30:00"
    assert event["available_quantity"] == 7

    ledger = client.get(ROOT, headers=admin).json()["items"]
    updated_stock = next(item for item in ledger if item["id"] == stock["id"])
    assert updated_stock["quantity"] == 12
    assert updated_stock["available_quantity"] == 7

    history = client.get(f"{ROOT}/outbounds?asset_id={stock['id']}", headers=admin)
    assert history.status_code == 200, history.text
    assert history.json()["total"] == 1
    assert history.json()["items"][0]["id"] == event["id"]
    assert history.json()["items"][0]["asset_code"] == stock["code"]

    insufficient = client.post(f"{ROOT}/{stock['id']}/outbound", headers=admin, json={
        "quantity": 8, "destination": "另一个客户现场", "shipped_at": "2026-10-01T15:00:00",
    })
    assert insufficient.status_code == 400
    with db() as session:
        asset = session.get(EquipmentAsset, stock["id"])
        assert asset.quantity == 12
        assert asset.available_quantity == 7
        assert session.query(EquipmentEvent).filter_by(asset_id=stock["id"], action="outbound").count() == 1


def test_inventory_contract_rejects_invalid_modes_and_protects_stock_history(client, db):
    admin, pm = login(client, "admin"), login(client, "pm")
    invalid_quantity = client.post(f"{ROOT}/inbound", headers=admin, json={
        "name": "错误单件数量", "items": [{
            "name": "笔记本", "category_code": "computer", "quantity": 2,
        }],
    })
    assert invalid_quantity.status_code == 400

    stock_batch = client.post(f"{ROOT}/inbound", headers=admin, json={
        "name": "出库权限测试库存", "inventory_managed": True, "items": [{
            "name": "光模块", "category_code": "network", "quantity": 4,
        }],
    })
    assert stock_batch.status_code == 200, stock_batch.text
    stock = stock_batch.json()["items"][0]
    assert client.post(f"{ROOT}/{stock['id']}/outbound", headers=pm, json={
        "quantity": 1, "customer": "客户", "shipped_at": "2026-10-01T15:00:00",
    }).status_code == 403
    assert client.post(f"{ROOT}/{stock['id']}/outbound", headers=admin, json={
        "quantity": 1, "shipped_at": "2026-10-01T15:00:00",
    }).status_code == 400
    delete_check = client.get(f"{ROOT}/{stock['id']}/delete-check", headers=admin)
    assert delete_check.status_code == 200
    assert delete_check.json()["allowed"] is False
    assert client.request("DELETE", f"{ROOT}/{stock['id']}", headers=admin, json={
        "confirm_code": stock["code"], "reason": "库存不可删除",
    }).status_code == 409


def test_inbound_does_not_treat_zero_batch_id_as_create_new_batch(client, db):
    admin = login(client, "admin")
    response = client.post(f"{ROOT}/inbound", headers=admin, json={
        "batch_id": 0, "name": "错误地提交的新批次", "items": [{"name": "设备", "category_code": "computer"}],
    })
    assert response.status_code == 400
    assert "请选择已有批次" in response.json()["detail"]
    with db() as session:
        assert session.query(EquipmentBatch).count() == 0
        assert session.query(EquipmentAsset).count() == 0


def test_inbound_can_reuse_batch_without_overwriting_metadata_and_inherits_acceptance(client, db):
    admin = login(client, "admin")
    created = client.post(f"{ROOT}/inbound", headers=admin, json={
        "name": "可复用设备批次", "source": "供应商送货", "supplier": "甲供应商",
        "acceptance_status": "pending_acceptance", "note": "原始备注",
        "items": [{"name": "第一台", "category_code": "computer"}],
    })
    assert created.status_code == 200, created.text
    batch_id = created.json()["batch_id"]

    appended = client.post(f"{ROOT}/inbound", headers=admin, json={
        "batch_id": batch_id, "items": [{"name": "第二台", "category_code": "computer"}],
    })
    assert appended.status_code == 200, appended.text
    assert appended.json()["batch_id"] == batch_id
    assert appended.json()["name"] == "可复用设备批次"
    assert appended.json()["asset_count"] == 2
    assert appended.json()["source"] == "供应商送货"
    assert appended.json()["supplier"] == "甲供应商"
    assert appended.json()["note"] == "原始备注"
    assert all(item["status"] == "pending_acceptance" for item in appended.json()["items"])

    with db() as session:
        assert session.query(EquipmentBatch).count() == 1
        assert session.query(EquipmentAsset).filter_by(batch_id=batch_id).count() == 2


def test_batch_admin_rename_search_and_delete_empty_batch(client, db):
    admin, pm = login(client, "admin"), login(client, "pm")
    created = client.post(f"{ROOT}/inbound", headers=admin, json={
        "name": "待整理批次", "items": [{"name": "临时设备", "category_code": "computer"}],
    })
    assert created.status_code == 200, created.text
    batch_id = created.json()["batch_id"]
    asset = created.json()["items"][0]

    assert client.patch(f"{ROOT}/batches/{batch_id}", headers=pm, json={"name": "无权更名"}).status_code == 403
    renamed = client.patch(f"{ROOT}/batches/{batch_id}", headers=admin, json={"name": "  设备批次·9月  "})
    assert renamed.status_code == 200, renamed.text
    assert renamed.json()["name"] == "设备批次·9月"
    assert client.get(f"{ROOT}/batches?keyword=设备批次", headers=pm).json()["items"][0]["asset_count"] == 1
    duplicate_name = client.post(f"{ROOT}/inbound", headers=admin, json={
        "name": "另一个批次", "items": [{"name": "另一台设备", "category_code": "computer"}],
    }).json()
    conflict = client.patch(f"{ROOT}/batches/{duplicate_name['batch_id']}", headers=admin,
                            json={"name": "设备批次·9月"})
    assert conflict.status_code == 400
    assert "批次名称已存在" in conflict.json()["detail"]

    assert client.delete(f"{ROOT}/batches/{batch_id}", headers=pm).status_code == 403
    assert client.delete(f"{ROOT}/batches/{batch_id}", headers=admin).status_code == 409
    removed = client.request("DELETE", f"{ROOT}/{asset['id']}", headers=admin,
                             json={"confirm_code": asset["code"], "reason": "清理测试设备"})
    assert removed.status_code == 200, removed.text

    empty = client.get(f"{ROOT}/batches?keyword=设备批次", headers=admin).json()["items"][0]
    assert empty["asset_count"] == 0
    deleted = client.delete(f"{ROOT}/batches/{batch_id}", headers=admin)
    assert deleted.status_code == 200, deleted.text
    assert deleted.json() == {"ok": True, "id": batch_id}
    with db() as session:
        assert session.get(EquipmentBatch, batch_id) is None
        assert session.query(AuditLog).filter_by(action="rename_equipment_batch", resource_id=str(batch_id)).count() == 1
        assert session.query(AuditLog).filter_by(action="delete_equipment_batch", resource_id=str(batch_id)).count() == 1


def test_batch_delete_rejects_orphan_event_reference(client, db):
    admin = login(client, "admin")
    batch = client.post(f"{ROOT}/inbound", headers=admin, json={
        "name": "遗留流水批次", "items": [{"name": "设备", "category_code": "computer"}],
    })
    assert batch.status_code == 200, batch.text
    batch_id = batch.json()["batch_id"]
    asset = batch.json()["items"][0]

    deleted = client.request("DELETE", f"{ROOT}/{asset['id']}", headers=admin,
                             json={"confirm_code": asset["code"], "reason": "清理测试设备"})
    assert deleted.status_code == 200, deleted.text
    standalone = _create(client, admin, name="批次外设备")
    with db() as session:
        session.execute(text("DROP TRIGGER equipment_events_immutable ON equipment_events"))
        session.add(EquipmentEvent(
            asset_id=standalone["id"], batch_id=batch_id, action="legacy", to_status="available",
            actor_id=session.query(User.id).filter_by(username="admin").scalar(), note="orphan guard",
        ))
        session.commit()
        session.execute(text(IMMUTABLE_TRIGGER))
        session.commit()

    removed = client.delete(f"{ROOT}/batches/{batch_id}", headers=admin)
    assert removed.status_code == 409
    assert "流水" in removed.json()["detail"]
    with db() as session:
        assert session.get(EquipmentBatch, batch_id) is not None
        assert session.query(EquipmentEvent).filter_by(batch_id=batch_id).count() == 1


def test_single_asset_creation_can_join_batch_and_inherits_acceptance_state(client):
    admin = login(client, "admin")
    created = client.post(f"{ROOT}/inbound", headers=admin, json={
        "name": "单件复用待验收", "acceptance_status": "pending_acceptance",
        "items": [{"name": "先录入", "category_code": "computer"}],
    })
    assert created.status_code == 200, created.text
    batch_id = created.json()["batch_id"]

    joined = client.post(ROOT, headers=admin, json={
        "name": "后录入单件", "category_code": "computer", "batch_id": batch_id,
    })
    assert joined.status_code == 200, joined.text
    assert joined.json()["batch_id"] == batch_id
    assert joined.json()["status"] == "pending_acceptance"


def test_batch_name_is_required_trimmed_and_case_insensitively_unique(client, db):
    admin = login(client, "admin")
    items = [{"name": "路由器", "category_code": "network"}]

    for invalid in ("", "   ", "x" * 129):
        response = client.post(f"{ROOT}/inbound", headers=admin, json={"name": invalid, "items": items})
        assert response.status_code == 400, response.text

    created = client.post(f"{ROOT}/inbound", headers=admin, json={"name": "  Lab Batch  ", "items": items})
    assert created.status_code == 200, created.text
    assert created.json()["name"] == "Lab Batch"

    duplicate = client.post(f"{ROOT}/inbound", headers=admin, json={"name": "lab batch", "items": items})
    assert duplicate.status_code == 400, duplicate.text
    assert "批次名称已存在" in duplicate.json()["detail"]
    with db() as session:
        assert session.query(EquipmentBatch).count() == 1


def test_export_neutralizes_formula_strings_in_csv_and_xlsx(client, db):
    admin = login(client, "admin")
    values = ["=1+1", "+SUM(A1:A2)", "-1+2", "@SUM(A1)", "  =1+1", "\t+1+1"]
    for i, value in enumerate(values):
        _create(client, admin, name=value, serial_number=f"FORMULA-{i}")
    # Normal asset creation trims these fields; place legacy/external values in
    # storage directly to exercise the export boundary's whitespace handling.
    with db() as session:
        for i, value in enumerate(values):
            asset = session.query(EquipmentAsset).filter_by(serial_number=f"FORMULA-{i}").one()
            asset.name = value
        session.commit()

    csv_response = client.get(f"{ROOT}/export?format=csv", headers=admin)
    assert csv_response.status_code == 200
    rows = list(csv.DictReader(io.StringIO(csv_response.content.decode("utf-8-sig"))))
    assert [row["设备名称"] for row in rows] == ["'" + value for value in values]

    xlsx_response = client.get(f"{ROOT}/export?format=xlsx", headers=admin)
    assert xlsx_response.status_code == 200
    from openpyxl import load_workbook
    sheet = load_workbook(io.BytesIO(xlsx_response.content), data_only=False).active
    assert [sheet.cell(row, 2).value for row in range(2, len(values) + 2)] == values
    assert all(sheet.cell(row, 2).data_type == "s" for row in range(2, len(values) + 2))


def test_pending_batch_metadata_is_readable_and_assets_cannot_be_claimed(client, db):
    admin, pm = login(client, "admin"), login(client, "pm")
    response = client.post(f"{ROOT}/inbound", headers=admin, json={
        "name": "待验收批次", "source": "供应商送货", "supplier": "星河设备有限公司", "purchase_order_no": "PO-2026-018",
        "invoice_no": "INV-2026-88", "acceptance_status": "pending_acceptance", "note": "待清点",
        "items": [{"name": "待验收电脑", "category_code": "computer", "serial_number": "PENDING-1"}],
    })
    assert response.status_code == 200, response.text
    result = response.json()
    asset = result["items"][0]
    assert asset["status"] == "pending_acceptance"
    assert client.post(f"{ROOT}/{asset['id']}/claim", headers=pm).status_code == 400
    assert client.post(f"{ROOT}/{asset['id']}/borrow", headers=pm, json={"due_date": "2099-01-01"}).status_code == 400
    batches = client.get(f"{ROOT}/batches", headers=pm)
    assert batches.status_code == 200
    batch = next(item for item in batches.json()["items"] if item["id"] == result["batch_id"])
    assert batch["source"] == "供应商送货"
    assert batch["supplier"] == "星河设备有限公司"
    assert batch["purchase_order_no"] == "PO-2026-018"
    assert batch["invoice_no"] == "INV-2026-88"
    assert batch["acceptance_status"] == "pending_acceptance"
    assert batch["note"] == "待清点"
    assert batch["created_by"]
    assert batch["created_at"]
    detail = client.get(f"{ROOT}/batches/{result['batch_id']}", headers=pm)
    assert detail.status_code == 200
    assert detail.json()["items"][0]["id"] == asset["id"]


def test_admin_acceptance_updates_batch_assets_events_and_audit_atomically(client, db, monkeypatch):
    admin, pm = login(client, "admin"), login(client, "pm")
    response = client.post(f"{ROOT}/inbound", headers=admin, json={
        "name": "验收审计回滚批次", "supplier": "Acme", "acceptance_status": "pending_acceptance", "items": [
            {"name": "电脑甲", "category_code": "computer"},
            {"name": "电脑乙", "category_code": "computer"},
        ],
    })
    assert response.status_code == 200, response.text
    batch_id = response.json()["batch_id"]
    assert client.post(f"{ROOT}/batches/{batch_id}/accept", headers=pm).status_code == 403
    accepted = client.post(f"{ROOT}/batches/{batch_id}/accept", headers=admin, json={"note": "验收合格"})
    assert accepted.status_code == 200, accepted.text
    assert accepted.json()["acceptance_status"] == "accepted"
    assert {item["status"] for item in accepted.json()["items"]} == {"available"}
    with db() as session:
        assert session.query(EquipmentEvent).filter_by(batch_id=batch_id, action="acceptance").count() == 2
        assert session.query(AuditLog).filter_by(resource="equipment_batches", resource_id=str(batch_id)).count() == 2

    pending = client.post(f"{ROOT}/inbound", headers=admin, json={
        "name": "二号验收批次", "acceptance_status": "pending_acceptance", "items": [{"name": "电脑丙", "category_code": "computer"}],
    }).json()
    second_batch_id = pending["batch_id"]
    import api.core.audit as audit
    def fail(*_args, **_kwargs):
        raise RuntimeError("audit unavailable")
    monkeypatch.setattr(audit, "audit_row", fail)
    failed = client.post(f"{ROOT}/batches/{second_batch_id}/accept", headers=admin)
    assert failed.status_code == 500
    with db() as session:
        assert session.get(EquipmentBatch, second_batch_id).acceptance_status == "pending_acceptance"
        second_asset = session.query(EquipmentAsset).filter_by(batch_id=second_batch_id).one()
        assert second_asset.status == "pending_acceptance"
        assert session.query(EquipmentEvent).filter_by(asset_id=second_asset.id).count() == 1


@pytest.mark.parametrize("extension,payload,content_type", [
    ("csv", "name,category_code\nImported,computer\nOther,monitor\n".encode(), "text/csv"),
    ("xlsx", None, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
])
def test_import_can_reuse_selected_batch_atomically(client, db, extension, payload, content_type):
    admin = login(client, "admin")
    pending = client.post(f"{ROOT}/inbound", headers=admin, json={
        "name": f"待导入批次-{extension}", "source": "既有来源", "supplier": "既有供应商",
        "acceptance_status": "pending_acceptance", "note": "不可覆盖",
        "items": [{"name": "先前设备", "category_code": "computer"}],
    })
    assert pending.status_code == 200, pending.text
    batch_id = pending.json()["batch_id"]
    if payload is None:
        book = Workbook()
        sheet = book.active
        sheet.append(["name", "category_code"])
        sheet.append(["Imported", "computer"])
        sheet.append(["Other", "monitor"])
        from io import BytesIO
        stream = BytesIO()
        book.save(stream)
        payload = stream.getvalue()
    response = client.post(f"{ROOT}/import", headers=admin, data={"batch_id": str(batch_id)},
                           files={"file": (f"assets.{extension}", payload, content_type)})
    assert response.status_code == 200, response.text
    assert response.json()["batch_id"] == batch_id
    assert response.json()["created"] == 2
    assert response.json()["batch_name"] == f"待导入批次-{extension}"
    with db() as session:
        batch = session.get(EquipmentBatch, batch_id)
        assert batch.source == "既有来源"
        assert batch.supplier == "既有供应商"
        assert batch.note == "不可覆盖"
        imported = session.query(EquipmentAsset).filter_by(batch_id=batch_id).order_by(EquipmentAsset.id).all()
        assert len(imported) == 3
        assert {asset.status for asset in imported} == {"pending_acceptance"}
        assert session.query(EquipmentEvent).filter_by(batch_id=batch_id).count() == 3


def test_import_creates_named_batch_after_row_prevalidation(client, db):
    admin = login(client, "admin")
    payload = "name,category_code\nImported,computer\n".encode()
    response = client.post(f"{ROOT}/import", headers=admin,
                           data={"batch_name": "台账导入批次", "acceptance_status": "accepted"},
                           files={"file": ("assets.csv", payload, "text/csv")})
    assert response.status_code == 200, response.text
    assert response.json()["batch_name"] == "台账导入批次"
    batches = client.get(f"{ROOT}/batches?keyword=台账导入批次", headers=admin).json()["items"]
    assert len(batches) == 1
    assert batches[0]["source"] == "导入 assets.csv"
    assert batches[0]["asset_count"] == 1


def test_import_prevalidation_is_atomic(client, db):
    admin = login(client, "admin")
    payload = b"name,category_code,serial_number\nGood,computer,S1\nBad,unknown,S2\n"
    response = client.post(f"{ROOT}/import", headers=admin, files={"file": ("assets.csv", payload, "text/csv")})
    assert response.status_code == 400, response.text
    assert response.json()["detail"]["errors"][0]["row"] == 3
    with db() as session:
        assert session.query(EquipmentAsset).count() == 0
        assert session.query(EquipmentBatch).count() == 0
        assert session.query(EquipmentEvent).count() == 0
    good = b"name,category_code,serial_number\nGood,computer,S1\nOther,monitor,S2\n"
    response = client.post(f"{ROOT}/import", headers=admin, data={"batch_name": "有效导入批次"},
                           files={"file": ("assets.csv", good, "text/csv")})
    assert response.status_code == 200, response.text
    assert response.json()["created"] == 2
    duplicate_import = client.post(f"{ROOT}/import", headers=admin, data={"batch_name": "重复导入批次"},
                                  files={"file": ("assets.csv", good, "text/csv")})
    assert duplicate_import.status_code == 400
    csv_response = client.get(f"{ROOT}/export?format=csv", headers=admin)
    assert csv_response.status_code == 200
    exported = list(csv.DictReader(io.StringIO(csv_response.content.decode("utf-8-sig"))))
    assert all(row["所属批次"].startswith("有效导入批次 · 批次 #") for row in exported)
    assert client.get(f"{ROOT}/export?format=xlsx", headers=admin).status_code == 200
    assert client.get(f"{ROOT}/export?format=csv", headers=login(client, "pm")).status_code == 403


def test_xlsx_import_and_invalid_rows(client, db):
    admin = login(client, "admin")
    book = Workbook()
    sheet = book.active
    sheet.append(["name", "category_code", "purchase_price"])
    sheet.append(["路由器", "network", 100])
    sheet.append(["无价设备", "other", "nonsense"])
    from io import BytesIO
    stream = BytesIO()
    book.save(stream)
    response = client.post(f"{ROOT}/import", headers=admin,
                           files={"file": ("assets.xlsx", stream.getvalue())})
    assert response.status_code == 400, response.text
    assert response.json()["detail"]["errors"][0]["row"] == 3
    with db() as session:
        assert session.query(EquipmentAsset).count() == 0


def test_corrupt_xlsx_returns_validation_error(client):
    admin = login(client, "admin")
    response = client.post(f"{ROOT}/import", headers=admin,
                           files={"file": ("assets.xlsx", b"broken workbook")})
    assert response.status_code == 400, response.text


def test_audit_failure_rolls_back_state_and_event(client, db, monkeypatch):
    admin = login(client, "admin")
    asset = _create(client, admin)
    import api.core.audit as audit
    def fail(*_args, **_kwargs):
        raise RuntimeError("audit unavailable")
    monkeypatch.setattr(audit, "audit_row", fail)
    response = client.post(f"{ROOT}/{asset['id']}/scrap", headers=admin)
    assert response.status_code == 500
    with db() as session:
        assert session.get(EquipmentAsset, asset["id"]).status == "available"
        assert session.query(EquipmentEvent).filter_by(asset_id=asset["id"]).count() == 1


def test_overdue_lists_only_borrowed_assets_past_due_date(client, db):
    admin, pm = login(client, "admin"), login(client, "pm")
    overdue_asset = _create(client, admin, name="逾期电脑")
    upcoming_asset = _create(client, admin, name="未逾期电脑")
    assert client.post(f"{ROOT}/{overdue_asset['id']}/borrow", headers=pm,
                       json={"due_date": "2099-01-01"}).status_code == 200
    assert client.post(f"{ROOT}/{upcoming_asset['id']}/borrow", headers=pm,
                       json={"due_date": "2099-01-01"}).status_code == 200
    with db() as session:
        session.get(EquipmentAsset, overdue_asset["id"]).due_date = date(2000, 1, 1)
        session.commit()

    response = client.get(f"{ROOT}/overdue", headers=pm)
    assert response.status_code == 200, response.text
    assert [item["id"] for item in response.json()["items"]] == [overdue_asset["id"]]


def test_category_admin_crud_and_referenced_category_cannot_be_deleted(client):
    admin = login(client, "admin")
    created = client.post(f"{ROOT}/categories", headers=admin,
                          json={"code": "camera", "label": "相机", "sort": 9})
    assert created.status_code == 200, created.text
    assert created.json() == {"code": "camera", "label": "相机", "sort": 9}
    updated = client.patch(f"{ROOT}/categories/camera", headers=admin,
                           json={"label": "摄影设备", "sort": 3})
    assert updated.status_code == 200, updated.text
    assert updated.json() == {"code": "camera", "label": "摄影设备", "sort": 3}
    assert client.delete(f"{ROOT}/categories/camera", headers=admin).json() == {"ok": True}

    assert client.post(f"{ROOT}/categories", headers=admin,
                       json={"code": "tablet", "label": "平板"}).status_code == 200
    assert client.post(ROOT, headers=admin,
                       json={"name": "平板设备", "category_code": "tablet"}).status_code == 200
    blocked = client.delete(f"{ROOT}/categories/tablet", headers=admin)
    assert blocked.status_code == 400


def test_statistics_monthly_inbound_and_outbound_count_equipment_events(client):
    admin, pm = login(client, "admin"), login(client, "pm")
    created = _create(client, admin, name="新建资产")
    inbound = client.post(f"{ROOT}/inbound", headers=admin, json={"name": "统计入库批次", "items": [
        {"name": "入库资产", "category_code": "computer"},
    ]})
    assert inbound.status_code == 200, inbound.text
    pending = client.post(f"{ROOT}/inbound", headers=admin, json={
        "name": "统计待验收批次", "acceptance_status": "pending_acceptance",
        "items": [{"name": "验收资产", "category_code": "monitor"}],
    })
    assert pending.status_code == 200, pending.text
    pending_id = pending.json()["batch_id"]
    assert client.post(f"{ROOT}/batches/{pending_id}/accept", headers=admin).status_code == 200
    assert client.post(f"{ROOT}/{created['id']}/claim", headers=pm).status_code == 200
    assert client.post(f"{ROOT}/{created['id']}/return", headers=admin).status_code == 200
    inbound_asset_id = inbound.json()["items"][0]["id"]
    assert client.post(f"{ROOT}/{inbound_asset_id}/borrow", headers=pm,
                       json={"due_date": "2099-01-01"}).status_code == 200

    response = client.get(f"{ROOT}/statistics", headers=pm)
    assert response.status_code == 200, response.text
    month = app_today().strftime("%Y-%m")
    summary = next(row for row in response.json()["by_month"] if row["month"] == month)
    assert summary == {"month": month, "inbound": 4, "outbound": 2}
