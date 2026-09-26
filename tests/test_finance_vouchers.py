"""P0-B regression tests for the finance voucher engine."""

from decimal import Decimal
import io
import time

from openpyxl import Workbook, load_workbook


def _auth(client, username, password):
    response = client.post(
        "/api/v1/auth/login", json={"username": username, "password": password}
    )
    assert response.status_code == 200
    return {"Authorization": f"Bearer {response.json()['access_token']}"}


def _book(client, headers, code="VOUCHER-TEST"):
    response = client.post(
        "/api/v1/finance/books", json={"code": code, "name": code}, headers=headers
    )
    assert response.status_code == 200
    return response.json()


def _scoped(headers, book):
    return {**headers, "X-Account-Book": book["code"]}


def _accounts(client, headers):
    rows = client.get("/api/v1/finance/accounts/tree", headers=headers).json()
    return {row["code"]: row for row in rows}


def test_admin_menu_exposes_finance_but_member_does_not(client):
    admin = _auth(client, "admin", "admin123456")
    member = _auth(client, "pm", "pm123456")

    admin_menus = client.get("/api/v1/meta/me/menus", headers=admin)
    member_menus = client.get("/api/v1/meta/me/menus", headers=member)

    assert admin_menus.status_code == member_menus.status_code == 200
    assert {item["href"] for item in admin_menus.json()["menus"]}.__contains__("/finance")
    assert "/finance" not in {item["href"] for item in member_menus.json()["menus"]}

    forbidden = client.get("/api/v1/finance/books", headers=member)
    assert forbidden.status_code == 403


def test_account_books_are_limited_to_cny(client, db):
    admin = _auth(client, "admin", "admin123456")
    accepted = client.post(
        "/api/v1/finance/books",
        headers=admin,
        json={"code": "BOOK-CNY-LOWER", "name": "人民币账套", "base_currency": "cny"},
    )
    assert accepted.status_code == 200
    assert accepted.json()["base_currency"] == "CNY"
    assert accepted.json()["currencies"] == ["CNY"]

    for code, currency in (("BOOK-USD", "USD"), ("BOOK-JPY", "JPY")):
        rejected = client.post(
            "/api/v1/finance/books",
            headers=admin,
            json={"code": code, "name": "非人民币账套", "base_currency": currency},
        )
        assert rejected.status_code == 400
        assert rejected.json()["detail"]["code"] == "F-BOOK-CURRENCY"
    with db() as session:
        from api.models import AccountBook
        assert session.query(AccountBook).filter(AccountBook.code.in_(["BOOK-USD", "BOOK-JPY"])).count() == 0


def test_voucher_rejects_unbalanced_entries_without_persisting(client):
    headers = _auth(client, "admin", "admin123456")
    book = _book(client, headers, "VOUCHER-UNBALANCED")
    scoped = _scoped(headers, book)
    accounts = _accounts(client, scoped)
    periods = client.get("/api/v1/finance/periods", headers=scoped).json()

    response = client.post(
        "/api/v1/finance/vouchers",
        headers=scoped,
        json={
            "period_id": periods[0]["id"],
            "voucher_no": "记-0001",
            "voucher_date": periods[0]["start_date"],
            "summary": "不平衡测试",
            "lines": [
                {"account_id": accounts["1405"]["id"], "debit": "10.00", "credit": "0", "summary": "采购"},
                {"account_id": accounts["6001"]["id"], "debit": "0", "credit": "9.99", "summary": "收入"},
            ],
        },
    )

    assert response.status_code == 400
    assert response.json()["detail"]["code"] == "F-UNBALANCED"
    assert client.get("/api/v1/finance/vouchers", headers=scoped).json()["total"] == 0


def test_voucher_rejects_internal_sources_and_unreviewed_posting(client):
    headers = _auth(client, "admin", "admin123456")
    book = _book(client, headers, "VOUCHER-SOURCE")
    scoped = _scoped(headers, book)
    accounts = _accounts(client, scoped)
    cash_account = next(row for row in accounts.values() if row["is_leaf"] and row["aux_flags"].get("cashflow"))
    period = client.get("/api/v1/finance/periods", headers=scoped).json()[0]
    payload = {
        "period_id": period["id"],
        "voucher_no": "记-0001",
        "voucher_date": period["start_date"],
        "source": "auto-close",
        "lines": [
            {"account_id": accounts["1405"]["id"], "debit": "1.00", "credit": "0", "summary": "库存"},
            {"account_id": accounts["6001"]["id"], "debit": "0", "credit": "1.00", "summary": "收入"},
        ],
    }
    rejected = client.post("/api/v1/finance/vouchers", headers=scoped, json=payload)
    assert rejected.status_code == 400
    assert rejected.json()["detail"]["code"] == "F-INVALID-VOUCHER"

    payload["source"] = "manual"
    created = client.post("/api/v1/finance/vouchers", headers=scoped, json=payload).json()
    assert client.post(f"/api/v1/finance/vouchers/{created['id']}/submit", headers=scoped).status_code == 200
    blocked = client.post(f"/api/v1/finance/vouchers/{created['id']}/post", headers=scoped)
    assert blocked.status_code == 400
    assert blocked.json()["detail"]["code"] == "F-SOD"


def test_voucher_lifecycle_posts_once_and_updates_balance(client):
    headers = _auth(client, "admin", "admin123456")
    book = _book(client, headers, "VOUCHER-LIFECYCLE")
    scoped = _scoped(headers, book)
    accounts = _accounts(client, scoped)
    periods = client.get("/api/v1/finance/periods", headers=scoped).json()
    payload = {
        "period_id": periods[0]["id"],
        "voucher_no": "记-0001",
        "voucher_date": periods[0]["start_date"],
        "summary": "销售收款",
        "lines": [
            {"account_id": accounts["1405"]["id"], "debit": "100.00", "credit": "0", "summary": "采购"},
            {"account_id": accounts["6001"]["id"], "debit": "0", "credit": "100.00", "summary": "收入"},
        ],
    }

    created = client.post("/api/v1/finance/vouchers", headers=scoped, json=payload)
    assert created.status_code == 200
    assert created.json()["status"] == "draft"

    voucher_id = created.json()["id"]
    submitted = client.post(f"/api/v1/finance/vouchers/{voucher_id}/submit", headers=scoped)
    assert submitted.status_code == 200
    assert submitted.json()["status"] == "pending"

    approved = client.post(f"/api/v1/finance/vouchers/{voucher_id}/approve", headers=scoped)
    assert approved.status_code == 200
    assert approved.json()["status"] == "pending"

    posted = client.post(f"/api/v1/finance/vouchers/{voucher_id}/post", headers=scoped)
    assert posted.status_code == 200
    assert posted.json()["status"] == "posted"

    duplicate = client.post(f"/api/v1/finance/vouchers/{voucher_id}/post", headers=scoped)
    assert duplicate.status_code == 400
    assert duplicate.json()["detail"]["code"] == "F-VOUCHER-STATE"

    balances = client.get(
        "/api/v1/finance/balances",
        headers=scoped,
        params={"period_id": periods[0]["id"]},
    )
    assert balances.status_code == 200
    rows = {row["account_code"]: row for row in balances.json()["items"]}
    assert Decimal(rows["1405"]["debit_total"]) == Decimal("100.00")
    assert Decimal(rows["6001"]["credit_total"]) == Decimal("100.00")


def test_posted_voucher_can_only_be_corrected_by_reversal(client):
    headers = _auth(client, "admin", "admin123456")
    book = _book(client, headers, "VOUCHER-REVERSAL")
    scoped = _scoped(headers, book)
    accounts = _accounts(client, scoped)
    cash_account = next(row for row in accounts.values() if row["is_leaf"] and row["aux_flags"].get("cashflow"))
    period = client.get("/api/v1/finance/periods", headers=scoped).json()[0]
    payload = {
        "period_id": period["id"],
        "voucher_no": "记-0001",
        "voucher_date": period["start_date"],
        "summary": "待冲销",
        "lines": [
            {"account_id": accounts["1405"]["id"], "debit": "20.00", "credit": "0", "summary": "采购"},
            {"account_id": accounts["6001"]["id"], "debit": "0", "credit": "20.00", "summary": "收入"},
        ],
    }
    created = client.post("/api/v1/finance/vouchers", headers=scoped, json=payload).json()
    vid = created["id"]
    assert client.post(f"/api/v1/finance/vouchers/{vid}/submit", headers=scoped).status_code == 200
    assert client.post(f"/api/v1/finance/vouchers/{vid}/approve", headers=scoped).status_code == 200
    assert client.post(f"/api/v1/finance/vouchers/{vid}/post", headers=scoped).status_code == 200

    edited = client.patch(
        f"/api/v1/finance/vouchers/{vid}", headers=scoped, json={"summary": "不允许直接改"}
    )
    assert edited.status_code == 400
    assert edited.json()["detail"]["code"] == "F-VOUCHER-POSTED"

    reversal = client.post(f"/api/v1/finance/vouchers/{vid}/reverse", headers=scoped)
    assert reversal.status_code == 200
    assert reversal.json()["status"] == "draft"
    assert reversal.json()["reversal_of"] == vid
    assert reversal.json()["lines"][0]["credit"] == "20.00"
    assert reversal.json()["lines"][0]["debit"] == "0.00"


def test_close_check_blocks_unreviewed_vouchers_and_close_round_trip(client):
    headers = _auth(client, "admin", "admin123456")
    book = _book(client, headers, "VOUCHER-CLOSE")
    scoped = _scoped(headers, book)
    accounts = _accounts(client, scoped)
    cash_account = next(row for row in accounts.values() if row["is_leaf"] and row["aux_flags"].get("cashflow"))
    period = client.get("/api/v1/finance/periods", headers=scoped).json()[0]
    payload = {
        "period_id": period["id"],
        "voucher_no": "记-0001",
        "voucher_date": period["start_date"],
        "summary": "结账检查",
        "lines": [
            {"account_id": accounts["1405"]["id"], "debit": "10.00", "credit": "0", "summary": "采购"},
            {"account_id": accounts["6001"]["id"], "debit": "0", "credit": "10.00", "summary": "收入"},
        ],
    }
    created = client.post("/api/v1/finance/vouchers", headers=scoped, json=payload).json()
    assert client.post(f"/api/v1/finance/vouchers/{created['id']}/submit", headers=scoped).status_code == 200

    blocked = client.post("/api/v1/finance/close/check", headers=scoped, params={"period_id": period["id"]})
    assert blocked.status_code == 200
    assert blocked.json()["ready"] is False
    assert any(item["code"] == "F-UNREVIEWED-VOUCHER" for item in blocked.json()["issues"])

    assert client.post(f"/api/v1/finance/vouchers/{created['id']}/approve", headers=scoped).status_code == 200
    assert client.post(f"/api/v1/finance/vouchers/{created['id']}/post", headers=scoped).status_code == 200
    closed = client.post("/api/v1/finance/close/run", headers={**scoped, "If-Match": str(period["version"])}, params={"period_id": period["id"]})
    assert closed.status_code == 200
    assert closed.json()["status"] == "closed"

    reopened = client.post("/api/v1/finance/close/reverse", headers={**scoped, "If-Match": "2"}, params={"period_id": period["id"], "confirm": "true"})
    assert reopened.status_code == 200
    assert reopened.json()["status"] == "open"
    vouchers = client.get("/api/v1/finance/vouchers", headers=scoped, params={"period_id": period["id"], "size": 100}).json()["items"]
    assert any(item["source"] == "auto-reopen" and item["status"] == "posted" for item in vouchers)


def test_posting_carries_previous_period_balance_forward(client):
    headers = _auth(client, "admin", "admin123456")
    book = _book(client, headers, "VOUCHER-CARRY-FORWARD")
    scoped = _scoped(headers, book)
    accounts = _accounts(client, scoped)
    periods = client.get("/api/v1/finance/periods", headers=scoped).json()
    first, second = periods[:2]

    def post(period, number, amount):
        payload = {
            "period_id": period["id"],
            "voucher_no": f"记-{number:04d}",
            "voucher_date": period["start_date"],
            "summary": "跨期余额",
            "lines": [
                {"account_id": accounts["1405"]["id"], "debit": str(amount), "credit": "0", "summary": "库存"},
                {"account_id": accounts["6001"]["id"], "debit": "0", "credit": str(amount), "summary": "收入"},
            ],
        }
        created = client.post("/api/v1/finance/vouchers", headers=scoped, json=payload).json()
        assert client.post(f"/api/v1/finance/vouchers/{created['id']}/submit", headers=scoped).status_code == 200
        assert client.post(f"/api/v1/finance/vouchers/{created['id']}/approve", headers=scoped).status_code == 200
        assert client.post(f"/api/v1/finance/vouchers/{created['id']}/post", headers=scoped).status_code == 200

    post(first, 1, "100.00")
    post(second, 1, "25.00")
    balances = client.get("/api/v1/finance/balances", headers=scoped, params={"period_id": second["id"]})
    assert balances.status_code == 200
    rows = {row["account_code"]: row for row in balances.json()["items"]}
    assert Decimal(rows["1405"]["debit_total"]) == Decimal("25.00")
    assert Decimal(rows["1405"]["ending_debit"]) == Decimal("125.00")
    assert Decimal(rows["6001"]["credit_total"]) == Decimal("25.00")
    assert Decimal(rows["6001"]["ending_credit"]) == Decimal("125.00")


def test_close_generates_profit_carry_forward_and_preserves_income_report(client):
    headers = _auth(client, "admin", "admin123456")
    book = _book(client, headers, "VOUCHER-AUTO-CLOSE")
    scoped = _scoped(headers, book)
    accounts = _accounts(client, scoped)
    period = client.get("/api/v1/finance/periods", headers=scoped).json()[0]
    payload = {
        "period_id": period["id"],
        "voucher_no": "记-0001",
        "voucher_date": period["start_date"],
        "summary": "结转测试收入",
        "lines": [
            {"account_id": accounts["1405"]["id"], "debit": "100.00", "credit": "0", "summary": "库存"},
            {"account_id": accounts["6001"]["id"], "debit": "0", "credit": "100.00", "summary": "收入"},
        ],
    }
    created = client.post("/api/v1/finance/vouchers", headers=scoped, json=payload).json()
    assert client.post(f"/api/v1/finance/vouchers/{created['id']}/submit", headers=scoped).status_code == 200
    assert client.post(f"/api/v1/finance/vouchers/{created['id']}/approve", headers=scoped).status_code == 200
    assert client.post(f"/api/v1/finance/vouchers/{created['id']}/post", headers=scoped).status_code == 200

    closed = client.post(
        "/api/v1/finance/close/run",
        headers={**scoped, "If-Match": str(period["version"])},
        params={"period_id": period["id"]},
    )
    assert closed.status_code == 200

    vouchers = client.get("/api/v1/finance/vouchers", headers=scoped, params={"period_id": period["id"], "size": 100}).json()["items"]
    auto = next(item for item in vouchers if item["source"] == "auto-close")
    assert auto["status"] == "posted"
    auto_lines = {line["account_code"]: line for line in auto["lines"]}
    assert auto_lines["6001"]["debit"] == "100.00"
    assert auto_lines["3103"]["credit"] == "100.00"

    income = client.get("/api/v1/finance/reports/income", headers=scoped, params={"period_id": period["id"]})
    assert income.status_code == 200
    assert Decimal(income.json()["totals"]["profit"]) == Decimal("100.00")
    balance_sheet = client.get("/api/v1/finance/reports/balance-sheet", headers=scoped, params={"period_id": period["id"]})
    assert balance_sheet.status_code == 200
    assert balance_sheet.json()["totals"]["balanced"] is True


def test_voucher_template_is_book_scoped_and_can_generate_draft(client):
    headers = _auth(client, "admin", "admin123456")
    book = _book(client, headers, "VOUCHER-TEMPLATE")
    scoped = _scoped(headers, book)
    accounts = _accounts(client, scoped)
    period = client.get("/api/v1/finance/periods", headers=scoped).json()[0]
    template = client.post(
        "/api/v1/finance/voucher-templates",
        headers=scoped,
        json={
            "code": "OFFICE",
            "name": "办公费模板",
            "lines": [
                {"account_id": accounts["1405"]["id"], "summary": "办公费", "debit": "100.00", "credit": "0", "aux_detail": {}},
                {"account_id": accounts["6001"]["id"], "summary": "办公费", "debit": "0", "credit": "100.00", "aux_detail": {}},
            ],
        },
    )
    assert template.status_code == 200
    assert template.json()["code"] == "OFFICE"
    listed = client.get("/api/v1/finance/voucher-templates", headers=scoped)
    assert listed.status_code == 200
    assert [row["code"] for row in listed.json()] == ["OFFICE"]

    generated = client.post(
        "/api/v1/finance/vouchers/from-template",
        headers=scoped,
        json={
            "template_id": template.json()["id"],
            "period_id": period["id"],
            "voucher_no": "记-0001",
            "voucher_date": period["start_date"],
            "summary": "一月办公费",
        },
    )
    assert generated.status_code == 200
    assert generated.json()["status"] == "draft"
    assert generated.json()["source"] == "template"
    assert generated.json()["lines"][0]["debit"] == "100.00"


def test_voucher_excel_import_is_atomic_and_reports_row_numbers(client):
    headers = _auth(client, "admin", "admin123456")
    book = _book(client, headers, "VOUCHER-IMPORT")
    scoped = _scoped(headers, book)
    accounts = _accounts(client, scoped)
    period = client.get("/api/v1/finance/periods", headers=scoped).json()[0]

    wb = Workbook()
    ws = wb.active
    ws.append(["凭证号", "凭证日期", "摘要", "科目编码", "分录摘要", "借方", "贷方", "现金流项目", "辅助核算JSON"])
    ws.append(["记-0001", period["start_date"], "导入成功", "1405", "库存", "12.00", "0", "", "{}"])
    ws.append(["记-0001", period["start_date"], "导入成功", "6001", "收入", "0", "12.00", "", "{}"])
    ws.append(["记-0002", period["start_date"], "导入失败", "NOPE", "错误科目", "1.00", "0", "", "{}"])
    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)

    imported = client.post(
        f"/api/v1/finance/vouchers/import?period_id={period['id']}",
        headers=scoped,
        files={"file": ("vouchers.xlsx", stream.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert imported.status_code == 400
    detail = imported.json()["detail"]
    assert detail["code"] == "F-IMPORT-INVALID"
    assert any(item["row"] == 4 for item in detail["details"])
    assert client.get("/api/v1/finance/vouchers", headers=scoped).json()["total"] == 0

    wb.remove(wb.active)
    ws = wb.create_sheet()
    ws.append(["凭证号", "凭证日期", "摘要", "科目编码", "分录摘要", "借方", "贷方", "现金流项目", "辅助核算JSON"])
    ws.append(["记-0001", period["start_date"], "导入成功", "1405", "库存", "12.00", "0", "", "{}"])
    ws.append(["记-0001", period["start_date"], "导入成功", "6001", "收入", "0", "12.00", "", "{}"])
    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)
    success = client.post(
        f"/api/v1/finance/vouchers/import?period_id={period['id']}",
        headers=scoped,
        files={"file": ("vouchers.xlsx", stream.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert success.status_code == 200
    assert success.json()["imported"] == 1


def test_voucher_excel_export_contains_current_book_vouchers(client):
    headers = _auth(client, "admin", "admin123456")
    book = _book(client, headers, "VOUCHER-EXPORT")
    scoped = _scoped(headers, book)
    accounts = _accounts(client, scoped)
    period = client.get("/api/v1/finance/periods", headers=scoped).json()[0]
    payload = {
        "period_id": period["id"],
        "voucher_no": "记-0001",
        "voucher_date": period["start_date"],
        "summary": "导出测试",
        "lines": [
            {"account_id": accounts["1405"]["id"], "debit": "2.00", "credit": "0", "summary": "库存"},
            {"account_id": accounts["6001"]["id"], "debit": "0", "credit": "2.00", "summary": "收入"},
        ],
    }
    created = client.post("/api/v1/finance/vouchers", headers=scoped, json=payload).json()
    assert client.post(f"/api/v1/finance/vouchers/{created['id']}/submit", headers=scoped).status_code == 200
    assert client.post(f"/api/v1/finance/vouchers/{created['id']}/approve", headers=scoped).status_code == 200
    assert client.post(f"/api/v1/finance/vouchers/{created['id']}/post", headers=scoped).status_code == 200
    exported = client.get(f"/api/v1/finance/vouchers/export?period_id={period['id']}", headers=scoped)
    assert exported.status_code == 200
    sheet = load_workbook(io.BytesIO(exported.content), read_only=True).active
    rows = list(sheet.iter_rows(values_only=True))
    assert rows[0][:5] == ("凭证号", "凭证日期", "摘要", "科目编码", "分录摘要")
    assert any(row[0] == "记-0001" for row in rows[1:])


def test_cash_transaction_creates_a_book_scoped_draft_voucher(client):
    headers = _auth(client, "admin", "admin123456")
    book = _book(client, headers, "CASH-TRANSACTION")
    scoped = _scoped(headers, book)
    accounts = _accounts(client, scoped)
    cash_account = next(row for row in accounts.values() if row["is_leaf"] and row["aux_flags"].get("cashflow"))
    period = client.get("/api/v1/finance/periods", headers=scoped).json()[0]

    created = client.post(
        "/api/v1/finance/cash-transactions",
        headers=scoped,
        json={
            "period_id": period["id"],
            "kind": "receipt",
            "transaction_date": period["start_date"],
            "amount": "88.00",
            "account_id": cash_account["id"],
            "offset_account_id": accounts["6001"]["id"],
            "summary": "客户收款",
            "counterparty": "客户A",
            "cashflow_item": "C01",
        },
    )
    assert created.status_code == 200, created.text
    body = created.json()
    assert body["kind"] == "receipt"
    assert body["voucher"]["status"] == "draft"
    assert body["voucher"]["source"] == "cash"
    assert body["voucher"]["lines"][0]["debit"] == "88.00"


def test_bank_statement_import_match_and_reconciliation_summary(client):
    headers = _auth(client, "admin", "admin123456")
    book = _book(client, headers, "BANK-RECONCILE")
    scoped = _scoped(headers, book)
    accounts = _accounts(client, scoped)
    cash_account = next(row for row in accounts.values() if row["is_leaf"] and row["aux_flags"].get("cashflow"))
    period = client.get("/api/v1/finance/periods", headers=scoped).json()[0]

    cash_response = client.post(
        "/api/v1/finance/cash-transactions",
        headers=scoped,
        json={
            "period_id": period["id"],
            "kind": "receipt",
            "transaction_date": period["start_date"],
            "amount": "50.00",
            "account_id": cash_account["id"],
            "offset_account_id": accounts["6001"]["id"],
            "summary": "银行收款",
            "cashflow_item": "C01",
        },
    )
    assert cash_response.status_code == 200, cash_response.text
    cash = cash_response.json()
    voucher_id = cash["voucher"]["id"]
    assert client.post(f"/api/v1/finance/vouchers/{voucher_id}/submit", headers=scoped).status_code == 200
    assert client.post(f"/api/v1/finance/vouchers/{voucher_id}/approve", headers=scoped).status_code == 200
    assert client.post(f"/api/v1/finance/vouchers/{voucher_id}/post", headers=scoped).status_code == 200

    wb = Workbook()
    ws = wb.active
    ws.append(["流水号", "交易日期", "方向", "金额", "银行账户", "摘要", "对方"])
    ws.append(["BANK-001", period["start_date"], "流入", "50.00", "100201", "银行收款", "客户A"])
    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)
    imported = client.post(
        f"/api/v1/finance/bank-statements/import?period_id={period['id']}",
        headers=scoped,
        files={"file": ("bank.xlsx", stream.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert imported.status_code == 200
    statement_id = imported.json()["items"][0]["id"]

    matched = client.post(
        f"/api/v1/finance/bank-statements/{statement_id}/match",
        headers=scoped,
        json={"voucher_id": voucher_id},
    )
    assert matched.status_code == 200
    assert matched.json()["status"] == "matched"

    summary = client.get(
        "/api/v1/finance/bank-reconciliation",
        headers=scoped,
        params={"period_id": period["id"]},
    )
    assert summary.status_code == 200
    assert summary.json()["unmatched_count"] == 0
    assert Decimal(summary.json()["bank_total"]) == Decimal("50.00")


def test_bank_statement_import_is_atomic_and_reports_row_number(client):
    headers = _auth(client, "admin", "admin123456")
    book = _book(client, headers, "BANK-IMPORT-ERROR")
    scoped = _scoped(headers, book)
    period = client.get("/api/v1/finance/periods", headers=scoped).json()[0]
    wb = Workbook()
    ws = wb.active
    ws.append(["流水号", "交易日期", "方向", "金额", "银行账户", "摘要", "对方"])
    ws.append(["BANK-OK", period["start_date"], "流入", "10.00", "100201", "正常", "A"])
    ws.append(["BANK-BAD", period["start_date"], "未知", "20.00", "100201", "错误", "B"])
    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)
    response = client.post(
        f"/api/v1/finance/bank-statements/import?period_id={period['id']}",
        headers=scoped,
        files={"file": ("bank.xlsx", stream.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert response.status_code == 400
    assert response.json()["detail"]["code"] == "F-IMPORT-INVALID"
    assert any(item["row"] == 3 for item in response.json()["detail"]["details"])
    assert client.get("/api/v1/finance/bank-statements", headers=scoped).json()["total"] == 0


def test_finance_book_roles_separate_read_entry_review_cash_and_close(client, db):
    admin = _auth(client, "admin", "admin123456")
    clerk = _auth(client, "pm", "pm123456")
    outsider = _auth(client, "outsider", "out123456")
    book = _book(client, admin, "ROLE-MATRIX")
    scoped_admin = _scoped(admin, book)
    accounts = _accounts(client, scoped_admin)
    period = client.get("/api/v1/finance/periods", headers=scoped_admin).json()[0]

    from api.models import AccountBookMember, User
    with db() as session:
        clerk_user = session.query(User).filter_by(username="pm").first()
        session.add(AccountBookMember(account_book_id=book["id"], user_id=clerk_user.id, role="finance_clerk"))
        session.commit()

    scoped_clerk = _scoped(clerk, book)
    read = client.get("/api/v1/finance/accounts/tree", headers=scoped_clerk)
    assert read.status_code == 200
    draft = client.post("/api/v1/finance/vouchers", headers=scoped_clerk, json={
        "period_id": period["id"], "voucher_no": "记-0001", "voucher_date": period["start_date"], "summary": "制单",
        "lines": [
            {"account_id": accounts["1405"]["id"], "debit": "1.00", "credit": "0", "summary": "库存"},
            {"account_id": accounts["6001"]["id"], "debit": "0", "credit": "1.00", "summary": "收入"},
        ],
    })
    assert draft.status_code == 200
    voucher_id = draft.json()["id"]
    assert client.post(f"/api/v1/finance/vouchers/{voucher_id}/submit", headers=scoped_clerk).status_code == 200
    assert client.post(f"/api/v1/finance/vouchers/{voucher_id}/approve", headers=scoped_clerk).status_code == 403
    assert client.post("/api/v1/finance/close/check", headers=scoped_clerk, params={"period_id": period["id"]}).status_code == 403

    assert client.get("/api/v1/finance/books", headers=outsider).status_code == 403


def test_fixed_asset_creation_and_next_period_straight_line_depreciation(client):
    headers = _auth(client, "admin", "admin123456")
    book = _book(client, headers, "FIXED-ASSET")
    scoped = _scoped(headers, book)
    accounts = _accounts(client, scoped)
    periods = client.get("/api/v1/finance/periods", headers=scoped).json()
    first, second = periods[:2]

    created = client.post(
        "/api/v1/finance/fixed-assets",
        headers=scoped,
        json={
            "code": "FA-001",
            "name": "测试设备",
            "category": "设备",
            "acquisition_date": first["start_date"],
            "original_cost": "1200.00",
            "residual_rate": "0.10",
            "useful_life_months": 12,
            "asset_account_id": accounts["1601"]["id"],
            "depreciation_account_id": accounts["1602"]["id"],
            "expense_account_id": accounts["6601"]["id"],
        },
    )
    assert created.status_code == 200, created.text
    assert created.json()["status"] == "active"
    assert created.json()["accumulated_depreciation"] == "0.00"

    same_month = client.post(
        f"/api/v1/finance/fixed-assets/depreciate?period_id={first['id']}",
        headers=scoped,
    )
    assert same_month.status_code == 200
    assert same_month.json()["depreciated"] == 0

    next_month = client.post(
        f"/api/v1/finance/fixed-assets/depreciate?period_id={second['id']}",
        headers=scoped,
    )
    assert next_month.status_code == 200, next_month.text
    assert next_month.json()["depreciated"] == 1
    assert next_month.json()["total_amount"] == "90.00"
    assert next_month.json()["vouchers"][0]["source"] == "auto-asset"
    listed = client.get("/api/v1/finance/fixed-assets", headers=scoped)
    assert listed.status_code == 200
    asset = listed.json()["items"][0]
    assert asset["accumulated_depreciation"] == "90.00"
    assert asset["net_book_value"] == "1110.00"


def _post_arap_voucher(client, headers, voucher_id):
    for action in ("submit", "approve", "post"):
        response = client.post(f"/api/v1/finance/vouchers/{voucher_id}/{action}", headers=headers)
        assert response.status_code == 200, response.text


def _create_arap_document(client, headers, period, party_id, account_id, offset_account_id, amount, kind, transaction_date=None, summary="测试往来单据"):
    endpoint = "/api/v1/finance/receivables" if kind == "receivable" else "/api/v1/finance/payables"
    account_key = "receivable_account_id" if kind == "receivable" else "payable_account_id"
    offset_key = "revenue_account_id" if kind == "receivable" else "expense_account_id"
    party_key = "customer_id" if kind == "receivable" else "vendor_id"
    response = client.post(endpoint, headers=headers, json={
        "period_id": period["id"], "transaction_date": transaction_date or period["start_date"],
        party_key: party_id, "amount": amount, account_key: account_id,
        offset_key: offset_account_id, "summary": summary,
    })
    assert response.status_code == 200, response.text
    # ARAP voucher numbers are generated from hundredths of a second.
    time.sleep(0.02)
    return response.json()


def test_receivable_collection_can_split_manual_allocations(client):
    headers = _auth(client, "admin", "admin123456")
    book = _book(client, headers, "AR-MANUAL-SPLIT")
    scoped = _scoped(headers, book)
    accounts = _accounts(client, scoped)
    period = client.get("/api/v1/finance/periods", headers=scoped).json()[0]
    customer = client.post("/api/v1/finance/auxiliary-items", headers=scoped, json={"kind": "customer", "code": "SPLIT-C01", "name": "拆分客户"}).json()
    other = client.post("/api/v1/finance/auxiliary-items", headers=scoped, json={"kind": "customer", "code": "SPLIT-C02", "name": "其他客户"}).json()
    first = _create_arap_document(client, scoped, period, customer["id"], accounts["11220101"]["id"], accounts["6001"]["id"], "60.00", "receivable")
    second = _create_arap_document(client, scoped, period, customer["id"], accounts["11220101"]["id"], accounts["6001"]["id"], "40.00", "receivable")
    outsider = _create_arap_document(client, scoped, period, other["id"], accounts["11220101"]["id"], accounts["6001"]["id"], "100.00", "receivable")
    for item in (first, second, outsider):
        _post_arap_voucher(client, scoped, item["voucher"]["id"])

    response = client.post("/api/v1/finance/receivable-collections", headers=scoped, json={
        "period_id": period["id"], "transaction_date": period["start_date"],
        "customer_id": customer["id"], "amount": "50.00",
        "cash_account_id": accounts["10020101"]["id"], "cashflow_item": "C01",
        "summary": "按单据指定收款",
        "allocations": [
            {"target_id": first["id"], "amount": "20.00"},
            {"target_id": second["id"], "amount": "30.00"},
        ],
    })

    assert response.status_code == 200, response.text
    assert response.json()["allocations"] == [
        {"receivable_id": first["id"], "amount": "20.00"},
        {"receivable_id": second["id"], "amount": "30.00"},
    ]
    items = {item["id"]: item for item in client.get("/api/v1/finance/receivables", headers=scoped).json()["items"]}
    assert items[first["id"]]["outstanding"] == "40.00"
    assert items[second["id"]]["outstanding"] == "10.00"
    assert items[outsider["id"]]["outstanding"] == "100.00"


def test_payable_payment_can_target_second_document_only(client):
    headers = _auth(client, "admin", "admin123456")
    book = _book(client, headers, "AP-MANUAL-TARGET")
    scoped = _scoped(headers, book)
    accounts = _accounts(client, scoped)
    period = client.get("/api/v1/finance/periods", headers=scoped).json()[0]
    vendor = client.post("/api/v1/finance/auxiliary-items", headers=scoped, json={"kind": "vendor", "code": "MANUAL-V01", "name": "指定供应商"}).json()
    first = _create_arap_document(client, scoped, period, vendor["id"], accounts["22020101"]["id"], accounts["6401"]["id"], "60.00", "payable")
    second = _create_arap_document(client, scoped, period, vendor["id"], accounts["22020101"]["id"], accounts["6401"]["id"], "40.00", "payable")
    _post_arap_voucher(client, scoped, first["voucher"]["id"])
    _post_arap_voucher(client, scoped, second["voucher"]["id"])

    response = client.post("/api/v1/finance/payable-payments", headers=scoped, json={
        "period_id": period["id"], "transaction_date": period["start_date"],
        "vendor_id": vendor["id"], "amount": "25.00",
        "cash_account_id": accounts["10020101"]["id"], "cashflow_item": "C02", "summary": "指定第二张付款",
        "allocations": [{"target_id": second["id"], "amount": "25.00"}],
    })

    assert response.status_code == 200, response.text
    assert response.json()["allocations"] == [{"payable_id": second["id"], "amount": "25.00"}]
    items = {item["id"]: item for item in client.get("/api/v1/finance/payables", headers=scoped).json()["items"]}
    assert items[first["id"]]["outstanding"] == "60.00"
    assert items[second["id"]]["outstanding"] == "15.00"


def test_manual_receivable_allocation_validation_is_atomic(client):
    headers = _auth(client, "admin", "admin123456")
    book = _book(client, headers, "AR-MANUAL-INVALID")
    scoped = _scoped(headers, book)
    accounts = _accounts(client, scoped)
    period = client.get("/api/v1/finance/periods", headers=scoped).json()[0]
    customer = client.post("/api/v1/finance/auxiliary-items", headers=scoped, json={"kind": "customer", "code": "INVALID-C01", "name": "核销校验客户"}).json()
    other = client.post("/api/v1/finance/auxiliary-items", headers=scoped, json={"kind": "customer", "code": "INVALID-C02", "name": "其他主体"}).json()
    first = _create_arap_document(client, scoped, period, customer["id"], accounts["11220101"]["id"], accounts["6001"]["id"], "60.00", "receivable")
    second = _create_arap_document(client, scoped, period, customer["id"], accounts["11220101"]["id"], accounts["6001"]["id"], "40.00", "receivable")
    foreign = _create_arap_document(client, scoped, period, other["id"], accounts["11220101"]["id"], accounts["6001"]["id"], "100.00", "receivable")
    for item in (first, second, foreign):
        _post_arap_voucher(client, scoped, item["voucher"]["id"])

    base = {
        "period_id": period["id"], "transaction_date": period["start_date"],
        "customer_id": customer["id"], "amount": "50.00",
        "cash_account_id": accounts["10020101"]["id"], "cashflow_item": "C01",
        "summary": "非法手工核销",
    }
    invalid = [
        {"allocations": [{"target_id": first["id"], "amount": "25.00"}, {"target_id": first["id"], "amount": "25.00"}]},
        {"allocations": [{"target_id": first["id"], "amount": "20.00"}]},
        {"allocations": [{"target_id": first["id"], "amount": "70.00"}]},
        {"allocations": [{"target_id": foreign["id"], "amount": "50.00"}]},
    ]
    for patch in invalid:
        before_vouchers = client.get("/api/v1/finance/vouchers", headers=scoped).json()["total"]
        before_items = {item["id"]: item["applied_amount"] for item in client.get("/api/v1/finance/receivables", headers=scoped).json()["items"]}
        response = client.post("/api/v1/finance/receivable-collections", headers=scoped, json={**base, **patch})
        assert response.status_code == 400, response.text
        assert isinstance(response.json().get("detail"), dict)
        assert response.json()["detail"].get("code")
        assert client.get("/api/v1/finance/vouchers", headers=scoped).json()["total"] == before_vouchers
        after_items = {item["id"]: item["applied_amount"] for item in client.get("/api/v1/finance/receivables", headers=scoped).json()["items"]}
        assert after_items == before_items


def test_manual_receivable_allocation_rejects_ineligible_document_states_and_dates(client):
    headers = _auth(client, "admin", "admin123456")
    book = _book(client, headers, "AR-MANUAL-ELIGIBILITY")
    scoped = _scoped(headers, book)
    accounts = _accounts(client, scoped)
    periods = client.get("/api/v1/finance/periods", headers=scoped).json()
    period = periods[0]
    customer = client.post("/api/v1/finance/auxiliary-items", headers=scoped, json={"kind": "customer", "code": "BOUNDARY-C01", "name": "边界客户"}).json()
    draft = _create_arap_document(client, scoped, period, customer["id"], accounts["11220101"]["id"], accounts["6001"]["id"], "20.00", "receivable", summary="草稿目标")
    future_date = period["end_date"]
    future = _create_arap_document(client, scoped, period, customer["id"], accounts["11220101"]["id"], accounts["6001"]["id"], "20.00", "receivable", transaction_date=future_date, summary="晚于收款日")
    _post_arap_voucher(client, scoped, future["voucher"]["id"])
    reversing = _create_arap_document(client, scoped, period, customer["id"], accounts["11220101"]["id"], accounts["6001"]["id"], "20.00", "receivable", summary="正在冲销目标")
    _post_arap_voucher(client, scoped, reversing["voucher"]["id"])
    reversal = client.post(f"/api/v1/finance/vouchers/{reversing['voucher']['id']}/reverse", headers=scoped)
    assert reversal.status_code == 200, reversal.text
    reversed_item = _create_arap_document(client, scoped, period, customer["id"], accounts["11220101"]["id"], accounts["6001"]["id"], "20.00", "receivable", summary="已冲销目标")
    _post_arap_voucher(client, scoped, reversed_item["voucher"]["id"])
    reversed_response = client.post(f"/api/v1/finance/vouchers/{reversed_item['voucher']['id']}/reverse", headers=scoped)
    assert reversed_response.status_code == 200, reversed_response.text
    _post_arap_voucher(client, scoped, reversed_response.json()["id"])

    for target, date in ((draft, period["start_date"]), (future, period["start_date"]), (reversing, period["start_date"]), (reversed_item, period["start_date"])):
        before_vouchers = client.get("/api/v1/finance/vouchers", headers=scoped).json()["total"]
        before_items = {item["id"]: item["applied_amount"] for item in client.get("/api/v1/finance/receivables", headers=scoped).json()["items"]}
        response = client.post("/api/v1/finance/receivable-collections", headers=scoped, json={
            "period_id": period["id"], "transaction_date": date, "customer_id": customer["id"],
            "amount": "5.00", "cash_account_id": accounts["10020101"]["id"], "cashflow_item": "C01",
            "summary": "目标资格边界", "allocations": [{"target_id": target["id"], "amount": "5.00"}],
        })
        assert response.status_code == 400, response.text
        assert isinstance(response.json().get("detail"), dict)
        assert response.json()["detail"].get("code")
        assert client.get("/api/v1/finance/vouchers", headers=scoped).json()["total"] == before_vouchers
        after_items = {item["id"]: item["applied_amount"] for item in client.get("/api/v1/finance/receivables", headers=scoped).json()["items"]}
        assert after_items == before_items


def test_fixed_asset_depreciation_is_idempotent_per_period(client):
    headers = _auth(client, "admin", "admin123456")
    book = _book(client, headers, "FIXED-ASSET-IDEMPOTENT")
    scoped = _scoped(headers, book)
    accounts = _accounts(client, scoped)
    periods = client.get("/api/v1/finance/periods", headers=scoped).json()
    asset = client.post(
        "/api/v1/finance/fixed-assets",
        headers=scoped,
        json={
            "code": "FA-002", "name": "服务器", "category": "设备",
            "acquisition_date": periods[0]["start_date"], "original_cost": "600.00",
            "residual_rate": "0", "useful_life_months": 6,
            "asset_account_id": accounts["1601"]["id"], "depreciation_account_id": accounts["1602"]["id"],
            "expense_account_id": accounts["6601"]["id"],
        },
    )
    assert asset.status_code == 200
    first = client.post(f"/api/v1/finance/fixed-assets/depreciate?period_id={periods[1]['id']}", headers=scoped)
    second = client.post(f"/api/v1/finance/fixed-assets/depreciate?period_id={periods[1]['id']}", headers=scoped)
    assert first.status_code == second.status_code == 200
    assert first.json()["depreciated"] == 1
    assert second.json()["depreciated"] == 0


def test_receivable_creation_and_fifo_collection_application(client):
    headers = _auth(client, "admin", "admin123456")
    book = _book(client, headers, "AR-AP")
    scoped = _scoped(headers, book)
    accounts = _accounts(client, scoped)
    period = client.get("/api/v1/finance/periods", headers=scoped).json()[0]
    customer = client.post("/api/v1/finance/auxiliary-items", headers=scoped, json={"kind": "customer", "code": "C001", "name": "客户A"})
    assert customer.status_code == 200
    customer_id = customer.json()["id"]

    receivable = client.post(
        "/api/v1/finance/receivables",
        headers=scoped,
        json={
            "period_id": period["id"], "transaction_date": period["start_date"], "customer_id": customer_id,
            "amount": "100.00", "receivable_account_id": accounts["11220101"]["id"],
            "revenue_account_id": accounts["6001"]["id"], "summary": "销售应收",
        },
    )
    assert receivable.status_code == 200, receivable.text
    receivable_id = receivable.json()["id"]
    voucher_id = receivable.json()["voucher"]["id"]
    assert client.post(f"/api/v1/finance/vouchers/{voucher_id}/submit", headers=scoped).status_code == 200
    assert client.post(f"/api/v1/finance/vouchers/{voucher_id}/approve", headers=scoped).status_code == 200
    assert client.post(f"/api/v1/finance/vouchers/{voucher_id}/post", headers=scoped).status_code == 200

    collection = client.post(
        "/api/v1/finance/receivable-collections",
        headers=scoped,
        json={"period_id": period["id"], "transaction_date": period["start_date"], "customer_id": customer_id, "amount": "40.00", "cash_account_id": accounts["10020101"]["id"], "cashflow_item": "C01", "summary": "收回部分货款"},
    )
    assert collection.status_code == 200, collection.text
    assert collection.json()["allocations"] == [{"receivable_id": receivable_id, "amount": "40.00"}]
    assert collection.json()["voucher"]["source"] == "cash"

    outstanding = client.get("/api/v1/finance/receivables", headers=scoped).json()["items"]
    assert outstanding[0]["outstanding"] == "60.00"


def test_payable_aging_buckets_and_over_collection_are_blocked(client):
    headers = _auth(client, "admin", "admin123456")
    book = _book(client, headers, "AR-AP-AGING")
    scoped = _scoped(headers, book)
    accounts = _accounts(client, scoped)
    period = client.get("/api/v1/finance/periods", headers=scoped).json()[0]
    vendor = client.post("/api/v1/finance/auxiliary-items", headers=scoped, json={"kind": "vendor", "code": "V001", "name": "供应商A"}).json()
    payable = client.post(
        "/api/v1/finance/payables",
        headers=scoped,
        json={"period_id": period["id"], "transaction_date": period["start_date"], "vendor_id": vendor["id"], "amount": "75.00", "payable_account_id": accounts["22020101"]["id"], "expense_account_id": accounts["6401"]["id"], "summary": "采购应付"},
    )
    assert payable.status_code == 200
    voucher_id = payable.json()["voucher"]["id"]
    assert client.post(f"/api/v1/finance/vouchers/{voucher_id}/submit", headers=scoped).status_code == 200
    assert client.post(f"/api/v1/finance/vouchers/{voucher_id}/approve", headers=scoped).status_code == 200
    assert client.post(f"/api/v1/finance/vouchers/{voucher_id}/post", headers=scoped).status_code == 200

    too_much = client.post("/api/v1/finance/payable-payments", headers=scoped, json={"period_id": period["id"], "transaction_date": period["start_date"], "vendor_id": vendor["id"], "amount": "80.00", "cash_account_id": accounts["10020101"]["id"], "summary": "超额付款"})
    assert too_much.status_code == 400
    assert too_much.json()["detail"]["code"] == "F-ARAP-OVER-APPLY"

    aging = client.get("/api/v1/finance/payables/aging", headers=scoped, params={"as_of": period["end_date"]})
    assert aging.status_code == 200
    assert aging.json()["total_outstanding"] == "75.00"
    assert aging.json()["buckets"]["0_30"] == "75.00"
    statement = client.get(
        "/api/v1/finance/payables/statement",
        headers=scoped,
        params={
            "party_id": vendor["id"], "account_id": accounts["22020101"]["id"],
            "from_date": period["start_date"], "to_date": period["end_date"],
        },
    )
    assert statement.status_code == 200
    assert statement.json()["period_increase"] == "75.00"
    assert statement.json()["closing_balance"] == "75.00"


def test_draft_arap_document_cannot_be_settled_before_posting(client):
    headers = _auth(client, "admin", "admin123456")
    book = _book(client, headers, "AR-AP-DRAFT-SETTLEMENT")
    scoped = _scoped(headers, book)
    accounts = _accounts(client, scoped)
    period = client.get("/api/v1/finance/periods", headers=scoped).json()[0]
    customer = client.post(
        "/api/v1/finance/auxiliary-items",
        headers=scoped,
        json={"kind": "customer", "code": "DRAFT-C01", "name": "草稿客户"},
    ).json()
    draft = client.post(
        "/api/v1/finance/receivables",
        headers=scoped,
        json={
            "period_id": period["id"], "transaction_date": period["start_date"],
            "customer_id": customer["id"], "amount": "100.00",
            "receivable_account_id": accounts["11220101"]["id"],
            "revenue_account_id": accounts["6001"]["id"], "summary": "尚未过账应收",
        },
    ).json()

    collection = client.post(
        "/api/v1/finance/receivable-collections",
        headers=scoped,
        json={
            "period_id": period["id"], "transaction_date": period["start_date"],
            "customer_id": customer["id"], "amount": "40.00",
            "cash_account_id": accounts["10020101"]["id"], "cashflow_item": "C01",
            "summary": "不能核销草稿应收",
        },
    )

    assert collection.status_code == 400
    assert collection.json()["detail"]["code"] == "F-ARAP-OVER-APPLY"
    listed = client.get("/api/v1/finance/receivables", headers=scoped).json()["items"]
    assert listed[0]["id"] == draft["id"]
    assert listed[0]["applied_amount"] == "0.00"


def test_arap_document_with_active_settlement_cannot_be_reversed(client):
    headers = _auth(client, "admin", "admin123456")
    book = _book(client, headers, "AR-AP-REVERSE-ALLOCATED")
    scoped = _scoped(headers, book)
    accounts = _accounts(client, scoped)
    period = client.get("/api/v1/finance/periods", headers=scoped).json()[0]
    customer = client.post(
        "/api/v1/finance/auxiliary-items",
        headers=scoped,
        json={"kind": "customer", "code": "REVERSE-C01", "name": "冲销客户"},
    ).json()
    receivable = client.post(
        "/api/v1/finance/receivables",
        headers=scoped,
        json={
            "period_id": period["id"], "transaction_date": period["start_date"],
            "customer_id": customer["id"], "amount": "100.00",
            "receivable_account_id": accounts["11220101"]["id"],
            "revenue_account_id": accounts["6001"]["id"], "summary": "有核销的应收",
        },
    ).json()
    invoice_id = receivable["voucher"]["id"]
    for action in ("submit", "approve", "post"):
        assert client.post(f"/api/v1/finance/vouchers/{invoice_id}/{action}", headers=scoped).status_code == 200
    collection = client.post(
        "/api/v1/finance/receivable-collections",
        headers=scoped,
        json={
            "period_id": period["id"], "transaction_date": period["start_date"],
            "customer_id": customer["id"], "amount": "40.00",
            "cash_account_id": accounts["10020101"]["id"], "cashflow_item": "C01",
            "summary": "已过账收款",
        },
    ).json()
    payment_id = collection["voucher"]["id"]
    for action in ("submit", "approve", "post"):
        assert client.post(f"/api/v1/finance/vouchers/{payment_id}/{action}", headers=scoped).status_code == 200

    reversed_invoice = client.post(f"/api/v1/finance/vouchers/{invoice_id}/reverse", headers=scoped)

    assert reversed_invoice.status_code == 409
    assert reversed_invoice.json()["detail"]["code"] == "F-ARAP-OPEN-ALLOCATIONS"

    payment_reversal = client.post(f"/api/v1/finance/vouchers/{payment_id}/reverse", headers=scoped)
    assert payment_reversal.status_code == 200
    payment_reversal_id = payment_reversal.json()["id"]
    for action in ("submit", "approve", "post"):
        assert client.post(f"/api/v1/finance/vouchers/{payment_reversal_id}/{action}", headers=scoped).status_code == 200

    invoice_reversal = client.post(f"/api/v1/finance/vouchers/{invoice_id}/reverse", headers=scoped)
    assert invoice_reversal.status_code == 200
    invoice_reversal_id = invoice_reversal.json()["id"]
    for action in ("submit", "approve", "post"):
        assert client.post(f"/api/v1/finance/vouchers/{invoice_reversal_id}/{action}", headers=scoped).status_code == 200

    aging = client.get(
        "/api/v1/finance/receivables/aging",
        headers=scoped,
        params={"as_of": period["end_date"], "party_id": customer["id"]},
    )
    assert aging.status_code == 200
    assert aging.json()["total_outstanding"] == "0.00"
    statement = client.get(
        "/api/v1/finance/receivables/statement",
        headers=scoped,
        params={"party_id": customer["id"], "from_date": period["start_date"], "to_date": period["end_date"]},
    )
    assert statement.status_code == 200
    assert statement.json()["period_increase"] == "140.00"
    assert statement.json()["period_decrease"] == "140.00"
    assert statement.json()["closing_balance"] == "0.00"
    assert len(statement.json()["items"]) == 4


def test_arap_document_cannot_be_settled_while_reversing_or_after_reversal(client):
    headers = _auth(client, "admin", "admin123456")
    book = _book(client, headers, "AR-AP-REVERSED-SETTLEMENT")
    scoped = _scoped(headers, book)
    accounts = _accounts(client, scoped)
    period = client.get("/api/v1/finance/periods", headers=scoped).json()[0]
    customer = client.post(
        "/api/v1/finance/auxiliary-items",
        headers=scoped,
        json={"kind": "customer", "code": "REVERSED-C01", "name": "冲销后客户"},
    ).json()
    receivable = client.post(
        "/api/v1/finance/receivables",
        headers=scoped,
        json={
            "period_id": period["id"], "transaction_date": period["start_date"],
            "customer_id": customer["id"], "amount": "100.00",
            "receivable_account_id": accounts["11220101"]["id"],
            "revenue_account_id": accounts["6001"]["id"], "summary": "准备冲销的应收",
        },
    ).json()
    invoice_id = receivable["voucher"]["id"]
    for action in ("submit", "approve", "post"):
        assert client.post(f"/api/v1/finance/vouchers/{invoice_id}/{action}", headers=scoped).status_code == 200

    reversal = client.post(f"/api/v1/finance/vouchers/{invoice_id}/reverse", headers=scoped)
    assert reversal.status_code == 200
    reversing_item = client.get("/api/v1/finance/receivables", headers=scoped).json()["items"][0]
    assert reversing_item["is_reversing"] is True
    payment_body = {
        "period_id": period["id"], "transaction_date": period["start_date"],
        "customer_id": customer["id"], "amount": "40.00",
        "cash_account_id": accounts["10020101"]["id"], "cashflow_item": "C01",
        "summary": "冲销期间不能收款",
    }
    while_reversing = client.post("/api/v1/finance/receivable-collections", headers=scoped, json=payment_body)
    assert while_reversing.status_code == 400
    assert while_reversing.json()["detail"]["code"] == "F-ARAP-OVER-APPLY"

    reversal_id = reversal.json()["id"]
    for action in ("submit", "approve", "post"):
        assert client.post(f"/api/v1/finance/vouchers/{reversal_id}/{action}", headers=scoped).status_code == 200
    after_reversal = client.post("/api/v1/finance/receivable-collections", headers=scoped, json=payment_body)
    assert after_reversal.status_code == 400
    assert after_reversal.json()["detail"]["code"] == "F-ARAP-OVER-APPLY"

    item = client.get("/api/v1/finance/receivables", headers=scoped).json()["items"][0]
    assert item["is_reversed"] is True
    assert item["is_reversing"] is False
    assert item["outstanding"] == "0.00"


def test_posting_arap_reversal_rechecks_active_allocations(client, db):
    from api.models import ARAPAllocation, ARAPTransaction

    headers = _auth(client, "admin", "admin123456")
    book = _book(client, headers, "AR-AP-REVERSE-POST-GUARD")
    scoped = _scoped(headers, book)
    accounts = _accounts(client, scoped)
    period = client.get("/api/v1/finance/periods", headers=scoped).json()[0]
    customer = client.post(
        "/api/v1/finance/auxiliary-items",
        headers=scoped,
        json={"kind": "customer", "code": "POST-GUARD-C01", "name": "过账保护客户"},
    ).json()
    receivable = client.post(
        "/api/v1/finance/receivables",
        headers=scoped,
        json={
            "period_id": period["id"], "transaction_date": period["start_date"],
            "customer_id": customer["id"], "amount": "100.00",
            "receivable_account_id": accounts["11220101"]["id"],
            "revenue_account_id": accounts["6001"]["id"], "summary": "并发保护应收",
        },
    ).json()
    invoice_id = receivable["voucher"]["id"]
    for action in ("submit", "approve", "post"):
        assert client.post(f"/api/v1/finance/vouchers/{invoice_id}/{action}", headers=scoped).status_code == 200
    payment_response = client.post(
        "/api/v1/finance/receivable-collections",
        headers=scoped,
        json={
            "period_id": period["id"], "transaction_date": period["start_date"],
            "customer_id": customer["id"], "amount": "10.00",
            "cash_account_id": accounts["10020101"]["id"], "summary": "并发保护收款",
            "cashflow_item": "C01",
        },
    )
    assert payment_response.status_code == 200, payment_response.text
    payment = payment_response.json()
    payment_id = payment["voucher"]["id"]
    for action in ("submit", "approve", "post"):
        assert client.post(f"/api/v1/finance/vouchers/{payment_id}/{action}", headers=scoped).status_code == 200
    with db() as session:
        session.query(ARAPAllocation).filter_by(
            account_book_id=book["id"], source_id=payment_id, target_id=receivable["id"],
        ).delete()
        target = session.query(ARAPTransaction).filter_by(id=receivable["id"]).one()
        target.applied_amount = Decimal("0.00")
        session.commit()

    reversal = client.post(f"/api/v1/finance/vouchers/{invoice_id}/reverse", headers=scoped)
    assert reversal.status_code == 200
    with db() as session:
        target = session.query(ARAPTransaction).filter_by(id=receivable["id"]).one()
        target.applied_amount = Decimal("10.00")
        session.add(ARAPAllocation(
            account_book_id=book["id"],
            source_id=payment_id,
            target_id=receivable["id"],
            amount=Decimal("10.00"),
        ))
        session.commit()

    reversal_id = reversal.json()["id"]
    for action in ("submit", "approve"):
        assert client.post(f"/api/v1/finance/vouchers/{reversal_id}/{action}", headers=scoped).status_code == 200
    blocked = client.post(f"/api/v1/finance/vouchers/{reversal_id}/post", headers=scoped)
    assert blocked.status_code == 409
    assert blocked.json()["detail"]["code"] == "F-ARAP-OPEN-ALLOCATIONS"

    with db() as session:
        session.query(ARAPAllocation).filter_by(account_book_id=book["id"], target_id=receivable["id"]).delete()
        target = session.query(ARAPTransaction).filter_by(id=receivable["id"]).one()
        target.applied_amount = Decimal("0.00")
        session.commit()
    posted = client.post(f"/api/v1/finance/vouchers/{reversal_id}/post", headers=scoped)
    assert posted.status_code == 200


def test_receivable_reports_use_posted_vouchers_and_statement_opening_balance(client):
    headers = _auth(client, "admin", "admin123456")
    book = _book(client, headers, "AR-AP-REPORTS")
    scoped = _scoped(headers, book)
    accounts = _accounts(client, scoped)
    periods = client.get("/api/v1/finance/periods", headers=scoped).json()
    january, february = periods[0], periods[1]
    customer = client.post(
        "/api/v1/finance/auxiliary-items",
        headers=scoped,
        json={"kind": "customer", "code": "REPORT-C01", "name": "账龄客户"},
    ).json()

    def receivable(period, amount, summary):
        return client.post(
            "/api/v1/finance/receivables",
            headers=scoped,
            json={
                "period_id": period["id"], "transaction_date": period["start_date"],
                "customer_id": customer["id"], "amount": amount,
                "receivable_account_id": accounts["11220101"]["id"],
                "revenue_account_id": accounts["6001"]["id"], "summary": summary,
            },
        ).json()

    def post_voucher(voucher_id):
        assert client.post(f"/api/v1/finance/vouchers/{voucher_id}/submit", headers=scoped).status_code == 200
        assert client.post(f"/api/v1/finance/vouchers/{voucher_id}/approve", headers=scoped).status_code == 200
        assert client.post(f"/api/v1/finance/vouchers/{voucher_id}/post", headers=scoped).status_code == 200

    posted_invoice = receivable(january, "100.00", "已过账应收")
    post_voucher(posted_invoice["voucher"]["id"])
    draft_invoice = receivable(january, "50.00", "草稿应收")
    listed_before_posting = client.get("/api/v1/finance/receivables", headers=scoped).json()["items"]
    draft_item = next(item for item in listed_before_posting if item["id"] == draft_invoice["id"])
    assert draft_item["voucher_status"] == "draft"

    draft_collection = client.post(
        "/api/v1/finance/receivable-collections",
        headers=scoped,
        json={
            "period_id": february["id"], "transaction_date": february["start_date"],
            "customer_id": customer["id"], "amount": "40.00",
            "cash_account_id": accounts["10020101"]["id"], "cashflow_item": "C01",
            "summary": "草稿收款",
        },
    ).json()

    aging = client.get(
        "/api/v1/finance/receivables/aging",
        headers=scoped,
        params={"as_of": february["end_date"], "party_id": customer["id"]},
    )
    assert aging.status_code == 200
    assert aging.json()["total_outstanding"] == "100.00"
    assert aging.json()["buckets"]["31_60"] == "100.00"

    post_voucher(draft_collection["voucher"]["id"])
    statement = client.get(
        "/api/v1/finance/receivables/statement",
        headers=scoped,
        params={
            "party_id": customer["id"], "account_id": accounts["11220101"]["id"],
            "from_date": february["start_date"], "to_date": february["end_date"],
        },
    )
    assert statement.status_code == 200
    assert statement.json()["opening_balance"] == "100.00"
    assert statement.json()["period_increase"] == "0.00"
    assert statement.json()["period_decrease"] == "40.00"
    assert statement.json()["closing_balance"] == "60.00"
    assert len(statement.json()["items"]) == 1
    assert statement.json()["items"][0]["voucher_id"] == draft_collection["voucher"]["id"]
    assert draft_invoice["voucher"]["status"] == "draft"
